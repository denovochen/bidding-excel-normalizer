"""跨 Sheet 项目关系、关系复核与标段降级回归。 @author denovochen"""
from __future__ import annotations

import csv
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from ledger_core.artifacts import publish
from ledger_core.normalize import build_outputs
from ledger_core.review import apply_answers, pending_questions, split_decisions
from ledger_core.workbook import inspect_workbooks, read_workbook


class CrossSheetRelationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def book(self, summary_rows, roster_rows):
        path = self.root / "input.xlsx"
        workbook = openpyxl.Workbook()
        summary = workbook.active
        summary.title = "结果汇总"
        for row in summary_rows:
            summary.append(row)
        roster = workbook.create_sheet("投标明细")
        for row in roster_rows:
            roster.append(row)
        workbook.save(path)
        workbook.close()
        return read_workbook(path)

    def plan(self, book):
        inspection = inspect_workbooks([book])
        plan = inspection["suggested_plan"]
        tables = [table for sheet in plan["sources"][0]["sheets"] for table in sheet.get("tables", [])]
        summary = next(table for table in tables if table["table_kind"] == "award_summary")
        roster = next(table for table in tables if table["table_kind"] == "bidder_roster")
        summary["award_completeness"] = {
            "status": "complete", "basis_type": "structural",
            "basis": "测试汇总区域完整列出最终中标结果",
        }
        roster["award_completeness"] = {
            "status": "partial", "basis_type": "none", "basis": "投标名单不含中标结果",
        }
        plan["relationships"] = [{
            "id": "summary_to_roster", "kind": "award_to_bidder_roster",
            "award_tables": [summary["table_id"]], "bidder_tables": [roster["table_id"]],
            "project_keys": ["project_code", "project_name"],
        }]
        return plan

    def test_exact_project_relation_moves_award_to_bidder_records(self):
        book = self.book(
            [["工程名称", "标段", "中标单位"], ["项目甲", "1标段", "甲公司"]],
            [["项目名称", "标段名称", "投标单位名称"],
             ["项目甲", "一标段", "甲公司"], ["项目甲", "一标段", "乙公司"]],
        )
        plan = self.plan(book)
        final, review, ledger = build_outputs([book], plan)
        self.assertEqual([(row["公司名称"], row["中标与否"]) for row in final], [("甲公司", "是"), ("乙公司", "否")])
        self.assertFalse(review)
        self.assertEqual(ledger["summary"]["relationship_count"], 1)
        self.assertEqual(ledger["summary"]["unresolved_relationship_count"], 0)
        self.assertEqual(ledger["relationships"][0]["basis"], "exact_project_key")
        self.assertTrue(all(record["sheet"] == "投标明细" for record in ledger["records"]))
        publish(self.root / "output", final, review, ledger)

    def test_multilevel_unit_name_uses_bidder_header_context(self):
        path = self.root / "headers.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["项目名称", "投标单位信息", None])
        sheet.append([None, "序号", "单位名称"])
        sheet.append(["项目甲", 1, "甲公司"])
        sheet.merge_cells("A1:A2")
        sheet.merge_cells("B1:C1")
        workbook.save(path)
        workbook.close()
        plan = inspect_workbooks([read_workbook(path)])["suggested_plan"]
        table = plan["sources"][0]["sheets"][0]["tables"][0]
        self.assertEqual(table["columns"]["bidder_name"], "C")
        self.assertEqual(table["table_kind"], "bidder_roster")

    def test_nonexact_project_requires_bounded_relation_selection(self):
        book = self.book(
            [["项目名称", "中标单位"], ["项目甲（财政补助）", "甲公司"]],
            [["工程名称", "投标企业名称"], ["项目甲", "甲公司"], ["项目甲", "乙公司"]],
        )
        plan = self.plan(book)
        final, review, draft = build_outputs([book], plan)
        questions, total, remaining = pending_questions(draft, {})
        self.assertEqual((total, remaining, len(questions)), (1, 1, 1))
        self.assertEqual(len(questions[0]["options"]), 2)
        self.assertFalse(questions[0]["allow_other"])
        state = {"decisions": {}}
        selected = questions[0]["options"][0]["value"]
        self.assertFalse(apply_answers(state, draft, {questions[0]["question_id"]: selected}))
        relation, award = split_decisions(state["decisions"])
        final, review, ledger = build_outputs([book], plan, award_resolutions=award,
                                              relation_resolutions=relation)
        self.assertEqual([(row["公司名称"], row["中标与否"]) for row in final], [("甲公司", "是"), ("乙公司", "否")])
        self.assertFalse(review)
        self.assertEqual(ledger["relationships"][0]["basis"], "user_selection")

    def test_deferred_relation_publishes_roster_and_standalone_review(self):
        book = self.book(
            [["项目名称", "中标单位"], ["项目甲（旧称）", "甲公司"]],
            [["项目名称", "投标单位名称"], ["项目甲", "甲公司"], ["项目甲", "乙公司"]],
        )
        plan = self.plan(book)
        _, _, draft = build_outputs([book], plan)
        question = pending_questions(draft, {})[0][0]
        state = {"decisions": {}}
        self.assertFalse(apply_answers(state, draft, {question["question_id"]: "unresolved"}))
        relation, award = split_decisions(state["decisions"])
        final, review, ledger = build_outputs([book], plan, award_resolutions=award,
                                              relation_resolutions=relation)
        self.assertEqual([row["中标与否"] for row in final], ["", ""])
        self.assertEqual(len(review), 1)
        self.assertEqual(ledger["summary"]["unresolved_relationship_count"], 1)
        publish(self.root / "deferred", final, review, ledger)

    def test_clear_multilot_routes_each_award_without_scope_issue(self):
        book = self.book(
            [["项目名称", "标段名称", "中标单位"],
             ["项目甲", "1标段", "甲公司"], ["项目甲", "2标段", "乙公司"]],
            [["项目名称", "标段", "投标单位名称"],
             ["项目甲", "一标段", "甲公司"], ["项目甲", "一标段", "乙公司"],
             ["项目甲", "二标段", "乙公司"], ["项目甲", "二标段", "丙公司"]],
        )
        plan = self.plan(book)
        final, review, ledger = build_outputs([book], plan)
        self.assertEqual([(row["标段名称"], row["公司名称"], row["中标与否"]) for row in final],
                         [("一标段", "甲公司", "是"), ("一标段", "乙公司", "否"),
                          ("二标段", "乙公司", "是"), ("二标段", "丙公司", "否")])
        self.assertNotIn("LOT_SCOPE_UNRESOLVED", [issue["code"] for issue in ledger["issues"]])
        self.assertFalse(review)

    def test_multilot_summary_with_single_roster_group_keeps_project_level_results_under_review(self):
        book = self.book(
            [["项目名称", "标段名称", "中标单位"],
             ["项目甲", "1标段", "甲公司"], ["项目甲", "2标段", "乙公司"]],
            [["项目名称", "投标单位名称"],
             ["项目甲", "甲公司"], ["项目甲", "乙公司"], ["项目甲", "丙公司"]],
        )
        plan = self.plan(book)
        final, review, ledger = build_outputs([book], plan)
        self.assertEqual([row["中标与否"] for row in final], ["是", "是", "否"])
        self.assertTrue(all(row["复核状态"] == "待复核" for row in final))
        self.assertEqual({issue["code"] for issue in ledger["issues"]}, {"LOT_SCOPE_UNRESOLVED"})
        self.assertEqual(len(review), 3)

    def test_cli_relation_review_precedes_award_review(self):
        book = self.book(
            [["项目名称", "中标单位"], ["项目甲（财政）", "甲建设工程有限公司"]],
            [["项目名称", "投标单位名称"], ["项目甲", "甲建设有限公司"], ["项目甲", "乙公司"]],
        )
        plan = self.plan(book)
        plan_path = self.root / "plan.json"
        plan_path.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
        output = self.root / "cli-output"
        first = subprocess.run([sys.executable, "-B", str(ROOT / "scripts/excel_ledger.py"), "run",
                                str(book.path), "--plan", str(plan_path), "--output", str(output)],
                               text=True, capture_output=True)
        self.assertEqual(first.returncode, 5, first.stderr)
        relation_reply = json.loads(first.stdout)
        self.assertEqual(relation_reply["kind"], "relationship_review_required")
        answers = self.root / "relation-answers.json"
        question = relation_reply["questions"][0]
        answers.write_text(json.dumps({question["question_id"]: question["options"][0]["value"]}), encoding="utf-8")
        second = subprocess.run([sys.executable, "-B", str(ROOT / "scripts/excel_ledger.py"), "resolve",
                                 "--state", relation_reply["state"], "--answers", str(answers)],
                                text=True, capture_output=True)
        self.assertEqual(second.returncode, 4, second.stderr)
        award_reply = json.loads(second.stdout)
        self.assertEqual(award_reply["kind"], "award_review_required")
        award_question = award_reply["questions"][0]
        answers.write_text(json.dumps({award_question["question_id"]: award_question["options"][0]["value"]}),
                           encoding="utf-8")
        final = subprocess.run([sys.executable, "-B", str(ROOT / "scripts/excel_ledger.py"), "resolve",
                                "--state", award_reply["state"], "--answers", str(answers)],
                               text=True, capture_output=True)
        self.assertEqual(final.returncode, 0, final.stderr)
        with (output / "final.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual([row["中标与否"] for row in rows], ["是", "否"])

    def test_compact_inspection_plan_patch_and_internal_cleanup(self):
        book = self.book(
            [["项目名称", "中标单位"], ["项目甲", "甲公司"]],
            [["项目名称", "投标单位名称"], ["项目甲", "甲公司"], ["项目甲", "乙公司"]],
        )
        output = self.root / "patched-output"
        initial = subprocess.run([sys.executable, "-B", str(ROOT / "scripts/excel_ledger.py"), "run",
                                  str(book.path), "--output", str(output)], text=True, capture_output=True)
        self.assertEqual(initial.returncode, 3, initial.stderr)
        reply = json.loads(initial.stdout)
        self.assertEqual(reply["kind"], "mapping_required")
        self.assertNotIn("suggested_plan", reply["inspection"])
        inspection_path = Path(reply["inspection"]["inspection_path"])
        self.assertTrue(inspection_path.is_file())
        candidate = reply["inspection"]["relationship_candidates"][0]
        stored = json.loads(inspection_path.read_text(encoding="utf-8"))
        tables = {table["table_id"]: table for source in stored["suggested_plan"]["sources"]
                  for sheet in source["sheets"] for table in sheet.get("tables", [])}
        summary_id = candidate["award_tables"][0]
        roster_id = candidate["bidder_tables"][0]
        patch = {
            "schema_version": 1,
            "table_updates": [
                {"table_id": summary_id, "set": {"award_completeness": {
                    "status": "complete", "basis_type": "structural", "basis": "测试汇总结果完整"}}},
                {"table_ids": [roster_id], "set": {"award_completeness": {
                    "status": "partial", "basis_type": "none", "basis": "测试明细不含结果"}}},
            ],
            "relationships": [{
                "id": "patched_relation", "kind": "award_to_bidder_roster",
                "award_tables": [summary_id], "bidder_tables": [roster_id],
                "project_keys": ["project_code", "project_name"],
            }],
        }
        patch_path = inspection_path.parent / "patch.json"
        plan_path = inspection_path.parent / "plan.json"
        patch_path.write_text(json.dumps(patch, ensure_ascii=False), encoding="utf-8")
        planned = subprocess.run([sys.executable, "-B", str(ROOT / "scripts/excel_ledger.py"), "plan",
                                  "--inspection", str(inspection_path), "--patch", str(patch_path),
                                  "--output", str(plan_path)], text=True, capture_output=True)
        self.assertEqual(planned.returncode, 0, planned.stderr)
        self.assertTrue(plan_path.is_file())
        result = subprocess.run([sys.executable, "-B", str(ROOT / "scripts/excel_ledger.py"), "run",
                                 str(book.path), "--plan", str(plan_path), "--output", str(output)],
                                text=True, capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertFalse(inspection_path.parent.exists())
        self.assertEqual({path.name for path in output.iterdir()}, {"final.csv", "review_queue.csv", "ledger.json"})


if __name__ == "__main__":
    unittest.main()
