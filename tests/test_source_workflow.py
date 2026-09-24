"""错误 plan、陌生布局与有界语义交接的端到端回归。 @author denovochen"""
from __future__ import annotations

import copy
import json
import re
import shlex
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from ledger_core import workflow
from ledger_core.artifacts import publish, validate_outputs
from ledger_core.contract import LedgerError, MappingRevisionRequired
from ledger_core.normalize import build_outputs
from ledger_core.workbook import inspect_workbooks, read_workbook, load_json


class SourceWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def book(self, rows, merges=(), name="陌生输入.xlsx", extra=None):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "随机页签"
        for row in rows:
            sheet.append(row)
        for bounds in merges:
            sheet.merge_cells(bounds)
        if extra:
            other = workbook.create_sheet("另一页")
            for row in extra:
                other.append(row)
        path = self.root / name
        workbook.save(path)
        workbook.close()
        return read_workbook(path)

    def plan(self, book):
        plan = inspect_workbooks([book])["suggested_plan"]
        for sheet in plan["sources"][0]["sheets"]:
            if sheet.get("tables"):
                sheet["action"] = "parse"
                for table in sheet["tables"]:
                    table["structure_warnings"] = []
                    for item in table["column_dispositions"]:
                        item.update(disposition="evidence", reason="合成来源辅助列")
        return plan

    def answer(self, question):
        answer = copy.deepcopy(question["answer_template"])
        answer["basis"] = "合成区域逐组列出完整最终中标结果"
        answer["award_completeness"] = "complete"
        return answer

    def resume(self, reply, answers):
        path = Path(reply["state"])
        return workflow.resume(path, load_json(path), answers)

    def test_platform_project_group_patch_is_rejected_before_dedup(self):
        book = self.book([["项目名称", "投标单位名称", "投标企业数量", "中标单位"],
                          ["项目甲", "甲公司", 2, "甲公司"], [None, "乙公司"],
                          [None, "甲公司", 2, "丙公司"], [None, "丙公司"]], ("A2:A5",))
        plan = self.plan(book)
        plan["sources"][0]["sheets"][0]["tables"][0]["group_mode"] = "project"
        with self.assertRaises(MappingRevisionRequired) as failure:
            build_outputs([book], plan)
        self.assertEqual(failure.exception.evidence["code"], "SOURCE_BLOCKS_COLLAPSED")

    def test_known_lot_column_cannot_be_disguised_as_evidence(self):
        book = self.book([["项目名称", "项目标段", "投标单位名称"],
                          ["项目甲", "一标段", "甲公司"], ["项目甲", "二标段", "甲公司"]])
        plan = self.plan(book)
        table = plan["sources"][0]["sheets"][0]["tables"][0]
        table["columns"].pop("lot_name")
        table["column_dispositions"].append({"column": "B", "disposition": "evidence", "reason": "平台错误判断"})
        table["group_mode"] = "project"
        with self.assertRaisesRegex(MappingRevisionRequired, "标段语义"):
            build_outputs([book], plan)

    def test_known_project_cannot_be_removed_to_bypass_project_coverage(self):
        book = self.book([["项目名称", "投标单位名称"], ["项目甲", "甲公司"]])
        plan = self.plan(book)
        table = plan["sources"][0]["sheets"][0]["tables"][0]
        table["columns"].pop("project_name")
        table["column_dispositions"].append({"column": "A", "disposition": "evidence", "reason": "错误降级"})
        table.update(project_mode="none", group_mode="source")
        with self.assertRaisesRegex(MappingRevisionRequired, "项目身份"):
            build_outputs([book], plan)

    def test_new_project_in_context_cannot_inherit_previous_project(self):
        book = self.book([["项目名称", "标段", "投标单位名称"],
                          ["甲建设项目", "一标段", "甲公司"], [None, "乙建设项目", "乙公司"]])
        plan = self.plan(book)
        table = plan["sources"][0]["sheets"][0]["tables"][0]
        table.update(group_mode="anchor", group_start_field="lot_name")
        with self.assertRaisesRegex(MappingRevisionRequired, "不能静默继承"):
            build_outputs([book], plan)

    def test_semantic_project_context_does_not_require_project_keyword(self):
        book = self.book([["项目名称", "标段", "投标单位名称"],
                          ["春雨田园", "一标段", "甲公司"], [None, "清溪灌区", "乙公司"]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        question = reply["questions"][0]
        answer = self.answer(question)
        answer["project_context_columns"] = ["B"]
        result = self.resume(reply, {question["question_id"]: answer})
        self.assertEqual(result["summary"]["project_count"], 2)
        ledger = load_json(self.root / "out" / "ledger.json")
        self.assertEqual([project["values"]["project_name"] for project in ledger["projects"]], ["春雨田园", "清溪灌区"])

    def test_named_lot_context_keeps_project_and_separate_participations(self):
        book = self.book([["项目名称", "标段", "投标单位名称"],
                          ["春雨田园", "东片", "甲公司"], [None, None, "乙公司"],
                          [None, "西片", "甲公司"], [None, None, "丙公司"]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        question = reply["questions"][0]
        answer = self.answer(question)
        answer["lot_context_columns"] = ["B"]
        result = self.resume(reply, {question["question_id"]: answer})
        self.assertEqual(result["summary"]["project_count"], 1)
        self.assertEqual(result["summary"]["group_count"], 2)
        self.assertEqual(result["summary"]["record_count"], 4)

    def test_unavailable_new_project_cannot_reuse_old_project(self):
        book = self.book([["项目名称", "投标单位名称"], ["项目甲", "甲公司"], ["=A2", "乙公司"]])
        plan = self.plan(book)
        plan["sources"][0]["sheets"][0]["tables"][0]["project_mode"] = "blocks"
        with self.assertRaisesRegex(MappingRevisionRequired, "缺少项目归属"):
            build_outputs([book], plan)

    def test_changed_project_year_with_missing_name_requires_review(self):
        book = self.book([["项目名称", "年度", "投标单位名称"], ["项目甲", "2022", "甲公司"],
                          [None, "2023", "乙公司"]])
        with self.assertRaisesRegex(MappingRevisionRequired, "上下文发生变化"):
            build_outputs([book], self.plan(book))

    def test_source_block_identity_does_not_depend_on_group_mode(self):
        book = self.book([["项目名称", "投标单位名称", "投标企业数量"],
                          ["项目甲", "甲公司", 2], [None, "乙公司"]])
        plan = self.plan(book)
        first = build_outputs([book], plan)[2]
        plan["sources"][0]["sheets"][0]["tables"][0]["group_mode"] = "project"
        second = build_outputs([book], plan)[2]
        self.assertEqual(first["source_blocks"], second["source_blocks"])
        self.assertNotEqual(first["groups"][0]["id"], second["groups"][0]["id"])

    def test_serial_restart_preserves_independent_participations(self):
        book = self.book([["项目名称", "投标序号", "投标单位名称"],
                          ["项目甲", 1, "甲公司"], ["项目甲", 2, "乙公司"],
                          ["项目甲", 1, "甲公司"], ["项目甲", 2, "丙公司"]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        result = self.resume(reply, {q["question_id"]: self.answer(q) for q in reply["questions"]})
        self.assertEqual(result["summary"]["record_count"], 4)
        self.assertEqual(result["summary"]["group_count"], 2)

    def test_truncated_candidate_group_cannot_recommend_wrong_winner(self):
        book = self.book([["项目名称", "标段", "投标单位名称", "中标单位"],
                          ["项目甲", "一标段", "甲建公司", "甲建设公司"],
                          [None, None, "乙公司"], [None, None, "甲建设公司"]], ("A2:A3", "B2:B4"))
        plan = self.plan(book)
        table = plan["sources"][0]["sheets"][0]["tables"][0]
        table.update(project_mode="merged", group_mode="anchor", group_start_field="lot_name")
        with self.assertRaises(MappingRevisionRequired):
            build_outputs([book], plan)
        table.update(project_mode="blocks", group_mode="source_blocks")
        final, _, ledger = build_outputs([book], plan)
        self.assertEqual(final[-1]["中标与否"], "是")
        self.assertEqual(ledger["groups"][0]["award_matches"][0]["candidate_count"], 3)

    def test_delete_source_record_cannot_pass_internal_consistency_only(self):
        book = self.book([["项目名称", "投标单位名称"], ["项目甲", "甲公司"], ["项目甲", "乙公司"]])
        final, review, ledger = build_outputs([book], self.plan(book))
        final.pop()
        ledger["records"].pop()
        ledger["summary"]["record_count"] = 1
        with self.assertRaisesRegex(LedgerError, "投标来源未完整输出"):
            publish(self.root / "bad", final, review, ledger)

    def test_reviewed_tail_of_source_block_stops_truncated_candidates(self):
        book = self.book([["项目名称", "标段", "投标单位名称", "中标单位"],
                          ["项目甲", "一标段", "甲建公司", "甲建设公司"],
                          [None, None, "乙公司"], [None, None, "甲建设公司"]], ("A2:A4", "B2:B4"))
        plan = self.plan(book)
        sheet = plan["sources"][0]["sheets"][0]
        sheet["tables"][0]["data_end_row"] = 3
        sheet["review_regions"] = [{"start": 4, "end": 4, "reason": "局部来源待确认"}]
        final, _, ledger = build_outputs([book], plan)
        self.assertTrue(all(not row["中标与否"] for row in final))
        self.assertIn("SOURCE_BLOCK_COVERAGE_GAP", {issue["code"] for issue in ledger["issues"]})
        self.assertTrue(all(not match["candidates"] for group in ledger["groups"] for match in group["award_matches"]))

    def test_splitting_tables_cannot_bypass_source_block_verification(self):
        book = self.book([["项目名称", "标段", "投标单位名称"],
                          ["项目甲", "一标段", "甲公司"], [None, None, "乙公司"], [None, None, "丙公司"]],
                         ("A2:A4", "B2:B4"))
        plan = self.plan(book)
        spec = plan["sources"][0]["sheets"][0]
        first = spec["tables"][0]
        first.pop("table_id")
        second = copy.deepcopy(first)
        first["data_end_row"] = 2
        second["data_start_row"] = 3
        spec["tables"] = [first, second]
        with self.assertRaisesRegex(MappingRevisionRequired, "不完整候选组"):
            build_outputs([book], plan)

    def test_legacy_plan_cannot_ignore_bidder_business_rows(self):
        book = self.book([["项目名称", "投标单位名称"], ["项目甲", "甲公司"], ["项目乙", "乙公司"]])
        plan = self.plan(book)
        spec = plan["sources"][0]["sheets"][0]
        spec["tables"][0]["data_end_row"] = 2
        spec["ignored_rows"] = [{"start": 3, "end": 3, "reason": "错误忽略"}]
        with self.assertRaisesRegex(MappingRevisionRequired, "潜在投标企业"):
            build_outputs([book], plan)

    def test_generic_semantics_compile_count_blocks_without_model_group_settings(self):
        book = self.book([["工程名称", "投标企业名称", "投标单位数量", "中标企业"],
                          ["同项目", "甲公司", 2, "甲公司"], [None, "乙公司"],
                          [None, "甲公司", 2, "丙公司"], [None, "丙公司"]], ("A2:A5",))
        reply = workflow.start([book.path], [book], [], self.root / "out")
        self.assertEqual(len(reply["questions"]), 1)
        question = reply["questions"][0]
        self.assertNotIn("group_mode", question["answer_template"])
        result = self.resume(reply, {question["question_id"]: self.answer(question)})
        self.assertEqual(result["kind"], "result")
        self.assertEqual(result["summary"]["record_count"], 4)
        self.assertEqual(result["summary"]["group_count"], 2)
        self.assertEqual(result["summary"]["review_record_count"], 0)
        ledger = load_json(self.root / "out" / "ledger.json")
        self.assertEqual(ledger["source_coverage"]["unaccounted_bidder_row_count"], 0)

    def test_project_only_workbook_stays_context_without_fabricating_companies(self):
        book = self.book([["项目名称", "项目编号"], ["项目甲", "P001"], ["项目乙", "P002"]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        question = reply["questions"][0]
        self.assertEqual(question["answer_template"]["record_layout"], "project_rows")
        result = self.resume(reply, {question["question_id"]: self.answer(question)})
        self.assertEqual(result["summary"]["record_count"], 2)
        self.assertEqual(result["summary"]["unique_company_count"], 0)
        self.assertEqual(result["summary"]["context_group_count"], 2)

    def test_permuted_columns_and_sheet_titles_do_not_change_participations(self):
        headers = ["项目名称", "标段名称", "投标企业名单", "中标单位"]
        rows = [["项目甲", "一标段", "甲公司、乙公司", "甲公司"],
                ["项目甲", "二标段", "甲公司、丙公司", "丙公司"]]
        for index, order in enumerate(((0, 1, 2, 3), (3, 2, 0, 1), (1, 3, 2, 0))):
            with self.subTest(order=order):
                book = self.book([[row[col] for col in order] for row in [headers, *rows]], name=f"不同名称{index}.xlsx")
                reply = workflow.start([book.path], [book], [], self.root / f"out{index}")
                result = self.resume(reply, {q["question_id"]: self.answer(q) for q in reply["questions"]})
                self.assertEqual(result["summary"]["record_count"], 4)
                self.assertEqual(result["summary"]["review_record_count"], 0)

    def test_repeated_headers_are_regions_and_same_structures_share_one_judgment(self):
        rows = [["项目名称", "投标单位名称"], ["项目甲", "甲公司"],
                ["项目名称", "投标单位名称"], ["项目乙", "甲公司"]]
        book = self.book(rows, extra=[["项目名称", "投标单位名称"], ["项目丙", "甲公司"]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        self.assertEqual(len(reply["questions"]), 1)
        self.assertEqual(len(reply["questions"][0]["regions"]), 3)
        result = self.resume(reply, {q["question_id"]: self.answer(q) for q in reply["questions"]})
        self.assertEqual(result["summary"]["record_count"], 3)

    def test_free_form_plan_parameters_are_rejected_and_review_is_bounded(self):
        book = self.book([["项目名称", "投标单位名称"], ["项目甲", "甲公司"]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        q = reply["questions"][0]
        wrong = {**self.answer(q), "group_mode": "project"}
        reply = self.resume(reply, {q["question_id"]: wrong})
        self.assertIn("不能设置分组参数", reply["questions"][0]["previous_error"]["message"])
        result = self.resume(reply, {q["question_id"]: wrong})
        self.assertEqual(result["kind"], "result")
        self.assertEqual(result["summary"]["record_count"], 0)
        self.assertEqual(result["summary"]["review_record_count"], 1)

    def test_cli_next_command_is_executable_with_quoted_paths(self):
        book = self.book([["项目名称", "投标单位名称"], ["项目甲", "甲公司"]], name="odd '$ name.xlsx")
        process = subprocess.run([sys.executable, "-B", str(ROOT / "scripts/excel_ledger.py"),
                                  "run", str(book.path), "--output", str(self.root / "out")], capture_output=True, text=True)
        self.assertEqual(process.returncode, 3, process.stderr)
        reply = json.loads(process.stdout)
        command = shlex.split(reply["next_command"])
        answers = Path(command[command.index("--answers") + 1])
        answers.write_text(json.dumps({q["question_id"]: self.answer(q) for q in reply["questions"]}), encoding="utf-8")
        process = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(process.returncode, 0, process.stderr)
        self.assertTrue(validate_outputs(self.root / "out")["validated"])

    def test_failed_context_block_is_isolated_without_discarding_other_projects(self):
        book = self.book([["项目名称", "标段", "投标单位名称"],
                          ["甲项目", "一标段", "甲公司"], [None, None, "乙公司"],
                          [None, "新的乙项目", "丙公司"], [None, None, "丁公司"],
                          ["丙项目", "一标段", "戊公司"]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        question = reply["questions"][0]
        answer = self.answer(question)
        reply = self.resume(reply, {question["question_id"]: answer})
        self.assertEqual(reply["kind"], "mapping_required")
        result = self.resume(reply, {question["question_id"]: answer})
        self.assertEqual(result["kind"], "result")
        self.assertEqual(result["summary"]["record_count"], 3)
        ledger = load_json(self.root / "out" / "ledger.json")
        issue = next(item for item in ledger["issues"] if item["code"] == "SOURCE_REGION_UNRESOLVED")
        self.assertEqual((issue["start_row"], issue["end_row"]), (4, 5))

    def test_unknown_headers_can_be_interpreted_without_new_adapter(self):
        book = self.book([["工程事项", "参标主体", "获选承接方"],
                          ["甲项目", "甲公司", "乙公司"], ["甲项目", "乙公司", None]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        question = reply["questions"][0]
        answer = self.answer(question)
        answer["columns"] = {"A": "project_name", "B": "bidder_name", "C": "award_name"}
        result = self.resume(reply, {question["question_id"]: answer})
        self.assertEqual(result["summary"]["record_count"], 2)
        self.assertEqual(result["summary"]["review_record_count"], 0)

    def test_unknown_headers_after_merged_title_need_no_file_adapter(self):
        book = self.book([["本期项目参建情况", None, None], ["事项", "参建主体", "最终承接方"],
                          ["甲项目", "甲公司", "乙公司"], ["甲项目", "乙公司", None]], ("A1:C1",))
        reply = workflow.start([book.path], [book], [], self.root / "out")
        question = reply["questions"][0]
        answer = self.answer(question)
        answer["header_rows"] = [2]
        answer["columns"] = {"A": "project_name", "B": "bidder_name", "C": "award_name"}
        result = self.resume(reply, {question["question_id"]: answer})
        self.assertEqual(result["summary"]["record_count"], 2)
        self.assertEqual(result["summary"]["review_record_count"], 0)

    def test_merged_business_headers_with_person_subheaders_are_not_companies(self):
        book = self.book([["项目名称", "中标企业名称", "负责人信息", None],
                          [None, None, "姓名", "身份证号"], ["项目甲", "甲公司", "张某", "证件"]],
                         ("A1:A2", "B1:B2", "C1:D1"))
        inspection = inspect_workbooks([book])
        self.assertEqual(inspection["suggested_plan"]["sources"][0]["sheets"][0]["tables"][0]["header_rows"], [1, 2])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        question = reply["questions"][0]
        answer = self.answer(question)
        answer["columns"] = {"A": "project_name", "B": "award_name", "C": "evidence", "D": "evidence"}
        result = self.resume(reply, {question["question_id"]: answer})
        self.assertEqual(result["summary"]["record_count"], 1)
        self.assertEqual(load_json(self.root / "out" / "ledger.json")["unique_companies"], ["甲公司"])

    def test_auxiliary_formula_issue_does_not_contaminate_other_bidders(self):
        book = self.book([["项目名称", "投标单位名称", "投标单位法人"],
                          ["项目甲", "甲公司", "=A1"], ["项目甲", "乙公司", "张某"]])
        final, review, ledger = build_outputs([book], self.plan(book))
        self.assertEqual([row["复核状态"] for row in final], ["通过", "通过"])
        self.assertEqual(len(review), 0)
        self.assertEqual(ledger["audit_warnings"][0]["source_cell"], "C2")
        self.assertEqual(ledger["audit_warnings"][0]["raw_value"], "=A1")
        self.assertTrue(ledger["groups"][0]["candidate_coverage"]["complete"])

    def test_total_row_does_not_create_a_fake_project(self):
        book = self.book([["项目名称", "中标单位"], ["合计", None], ["项目甲", "甲公司"]])
        final, _, ledger = build_outputs([book], self.plan(book))
        self.assertEqual(len(final), 1)
        self.assertEqual(ledger["summary"]["project_count"], 1)

    def test_unlabelled_amount_total_does_not_create_empty_business_record(self):
        book = self.book([["项目名称", "中标单位", "中标金额"],
                          [None, None, "=SUM(C3:C4)"], ["项目甲", "甲公司", 100]])
        final, review, ledger = build_outputs([book], self.plan(book))
        self.assertEqual(len(final), 1)
        self.assertEqual(final[0]["公司名称"], "甲公司")
        self.assertFalse(review)
        self.assertEqual(ledger["summary"]["group_count"], 1)

    def test_project_and_count_only_block_head_belongs_to_following_bidders(self):
        book = self.book([["项目名称", "投标单位名称", "投标企业数量", "中标单位"],
                          ["项目甲", None, 2, "乙公司"], [None, "甲公司"], [None, "乙公司"]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        result = self.resume(reply, {q["question_id"]: self.answer(q) for q in reply["questions"]})
        self.assertEqual(result["summary"]["record_count"], 2)
        self.assertEqual(result["summary"]["group_count"], 1)
        self.assertEqual(result["summary"]["review_record_count"], 0)

    def test_long_mapping_rationale_is_not_repeated_for_every_source_cell(self):
        sizes = []
        for count in (10, 100):
            book = self.book([["项目名称", "投标单位名称", "辅助信息"],
                              *[["项目甲", f"企业{index}公司", "原始值"] for index in range(count)]], name=f"length{count}.xlsx")
            plan = self.plan(book)
            plan["sources"][0]["sheets"][0]["tables"][0]["column_dispositions"][0]["reason"] = "来源语义说明" * 1000
            result = build_outputs([book], plan)
            sizes.append(len(json.dumps(result[2], ensure_ascii=False).encode("utf-8")))
            publish(self.root / f"size{count}", *result)
        self.assertLess(sizes[1] - sizes[0], 400000)

    def test_parallel_duplicate_business_columns_cannot_be_silently_dropped(self):
        book = self.book([["项目名称", "投标单位名称", "项目名称", "投标单位名称"],
                          ["甲项目", "甲公司", "乙项目", "乙公司"]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        question = reply["questions"][0]
        answer = self.answer(question)
        answer["columns"] = {"A": "project_name", "B": "bidder_name", "C": "evidence", "D": "evidence"}
        reply = self.resume(reply, {question["question_id"]: answer})
        self.assertIn("并排重复业务列", reply["questions"][0]["previous_error"]["message"])

    def test_wide_headers_have_bounded_handoff_and_column_evidence_pages(self):
        book = self.book([["项目名称", "投标单位名称", *[f"很长的自定义辅助表头{index}" * 5 for index in range(150)]],
                          ["甲项目", "甲公司", *["x" * 200 for _ in range(150)]]])
        reply = workflow.start([book.path], [book], [], self.root / "out")
        self.assertLess(len(json.dumps(reply, ensure_ascii=False)), workflow.MAX_RESPONSE_CHARS)
        question = reply["questions"][0]
        state_path = Path(reply["state"])
        evidence = workflow.inspect_region(state_path, load_json(state_path), question["regions"][0]["id"], "1:2", "Q:AF")
        self.assertEqual(len(evidence["columns"]), 16)
        self.assertTrue(all(int(re.search(r"\d+$", cell)[0]) <= 2 for row in evidence["rows"] for cell in row))


if __name__ == "__main__":
    unittest.main()
