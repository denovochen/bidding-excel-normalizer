"""稀疏 XLSX 范围、合并继承及文件级失败隔离回归。 @author denovochen"""
from __future__ import annotations

import csv
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.table import Table

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from ledger_core import workbook as reader
from ledger_core.contract import LedgerError, RecoverableWorkbookError
from ledger_core.normalize import build_outputs

Q = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
FILL = PatternFill("solid", fgColor="FFFF00")


class SparseWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def save(self, configure, name="input.xlsx"):
        book = openpyxl.Workbook()
        book.active.title = "台账"
        configure(book.active)
        path = self.directory / name
        book.save(path)
        book.close()
        return path

    def rewrite(self, path, part, change):
        # 只修改合成夹具的 XML，覆盖 Excel 库不会自然生成的合法边界。
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        root = ET.fromstring(entries[part])
        change(root)
        entries[part] = ET.tostring(root)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)

    def cli(self, *args):
        result = subprocess.run([sys.executable, "-B", str(ROOT / "scripts/excel_ledger.py"), *map(str, args)],
                                capture_output=True, text=True)
        return result, json.loads(result.stdout) if result.stdout else None

    def test_style_only_xfd_does_not_move_or_drop_values(self):
        def configure(sheet):
            sheet["A1"] = "项目名称"
            sheet["D2"] = "项目甲"
            sheet["M3"] = 0
            sheet["XFD515"].fill = FILL
        path = self.save(configure)
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        sheet = reader.read_workbook(path).sheets[0]
        self.assertEqual(set(sheet.cells), {(1, 1), (2, 4), (3, 13)})
        self.assertEqual((sheet.max_row, sheet.max_col), (3, 13))
        self.assertEqual(sheet.raw(3, 13).value, 0)
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_far_content_types_fail_before_openpyxl_load(self):
        for value in ("业务值", 0, False, '=""', "#REF!"):
            with self.subTest(value=value):
                path = self.save(lambda s: s.__setitem__("XFD1", value))
                with patch.object(openpyxl, "load_workbook") as load:
                    with self.assertRaisesRegex(RecoverableWorkbookError, "Sheet=台账.*XFD1"):
                        reader.read_workbook(path)
                    load.assert_not_called()

    def test_empty_formula_and_error_payloads_are_not_blank_styles(self):
        for kind in ("f", "e"):
            with self.subTest(kind=kind):
                path = self.save(lambda s: s.__setitem__("XFD1", "=1" if kind == "f" else "#N/A"))
                def change(root):
                    cell = root.find(".//" + Q + "c")
                    for child in cell:
                        child.text = None
                self.rewrite(path, "xl/worksheets/sheet1.xml", change)
                with self.assertRaises(RecoverableWorkbookError):
                    reader.read_workbook(path)

    def test_row_and_rectangle_limits_still_reject_real_content(self):
        for address in ("A100001", "T30000"):
            with self.subTest(address=address):
                path = self.save(lambda s: s.__setitem__(address, "值"))
                with self.assertRaisesRegex(RecoverableWorkbookError, address):
                    reader.read_workbook(path)

    def test_empty_far_rows_are_not_scanned(self):
        def configure(sheet):
            sheet["A1"] = "保留"
            sheet["XFD1048576"].fill = FILL
        sheet = reader.read_workbook(self.save(configure)).sheets[0]
        self.assertEqual(set(sheet.cells), {(1, 1)})

    def test_vertical_merge_keeps_project_inheritance(self):
        def configure(sheet):
            for row in (("项目名称", "投标单位名称", "中标单位"), ("项目甲", "甲公司", "甲公司"),
                        (None, "乙公司", None)):
                sheet.append(row)
            sheet.merge_cells("A2:A3")
            sheet["XFD1"].fill = FILL
        book = reader.read_workbook(self.save(configure))
        sheet = book.sheets[0]
        cell, row, col = sheet.resolved(3, 1)
        self.assertEqual((cell.value, row, col), ("项目甲", 2, 1))
        plan = reader.inspect_workbooks([book])["suggested_plan"]
        final, _, _ = build_outputs([book], plan)
        self.assertEqual([r["项目名称"] for r in final], ["项目甲", "项目甲"])

    def test_horizontal_multilevel_headers_keep_inherited_text(self):
        def configure(sheet):
            for row in (("项目基本情况", None, "招投标情况", None),
                        ("项目名称", "标段", "投标企业名单", "中标单位"),
                        ("项目甲", "一标", "甲公司、乙公司", "甲公司")):
                sheet.append(row)
            sheet.merge_cells("A1:B1")
            sheet.merge_cells("C1:D1")
            sheet["XFD1"].fill = FILL
        book = reader.read_workbook(self.save(configure))
        self.assertEqual(reader._header_text(book.sheets[0], [1, 2], 4), "招投标情况 / 中标单位")
        plan = reader.inspect_workbooks([book])["suggested_plan"]
        spec = plan["sources"][0]["sheets"][0]
        spec["tables"][0]["header_rows"] = [1, 2]
        spec["ignored_rows"] = []
        self.assertEqual(len(build_outputs([book], plan)[0]), 2)

    def test_merge_can_extend_beyond_last_populated_column(self):
        def configure(sheet):
            sheet["M4"] = "合并内容"
            sheet.merge_cells("M4:N4")
            sheet.merge_cells("XFC10:XFD20")
        sheet = reader.read_workbook(self.save(configure)).sheets[0]
        self.assertEqual(sheet.merges, [(4, 13, 4, 14)])
        self.assertEqual(sheet.resolved(4, 14)[0].value, "合并内容")
        self.assertNotIn((4, 14), sheet.cells)

    def test_blank_intersection_is_retained_without_recursive_expansion(self):
        def configure(sheet):
            sheet["A1"] = "首"
            sheet["B4"] = "尾"
            sheet.merge_cells("B2:C2")
            sheet.merge_cells("C3:D3")
            sheet.merge_cells("XFC10:XFD20")
        sheet = reader.read_workbook(self.save(configure)).sheets[0]
        self.assertEqual(sheet.merges, [(2, 2, 2, 3)])

    def test_meaningful_and_intersecting_oversized_merges_fail(self):
        for value, area in (("内容", "A1:XFD1"), ('=""', "A1:XFD1"), (None, "A2:XFD2")):
            with self.subTest(value=value, area=area):
                def configure(sheet):
                    sheet["A1"] = value
                    sheet["B1"] = "上边界"
                    sheet["A3"] = "范围内内容"
                    sheet.merge_cells(area)
                path = self.save(configure)
                with self.assertRaisesRegex(RecoverableWorkbookError, "合并=.*XFD"):
                    reader.read_workbook(path)

    def test_remote_merge_with_content_is_not_filtered_out(self):
        def configure(sheet):
            sheet["A1"] = "正常内容"
            sheet["XFC10"] = "远端合并内容"
            sheet.merge_cells("XFC10:XFD20")
        with self.assertRaisesRegex(RecoverableWorkbookError, "XFC10"):
            reader.read_workbook(self.save(configure))

    def test_merge_cannot_mask_serialized_nonanchor_content(self):
        def configure(sheet):
            sheet["A1"] = "合并锚点"
            sheet.merge_cells("A1:A2")
            sheet["B2"] = "正常数据"
        path = self.save(configure)
        def change(root):
            row = root.find(f".//{Q}row[@r='2']")
            cell = ET.Element(Q + "c", {"r": "A2", "t": "n"})
            ET.SubElement(cell, Q + "v").text = "123"
            row.insert(0, cell)
        self.rewrite(path, "xl/worksheets/sheet1.xml", change)
        with self.assertRaisesRegex(RecoverableWorkbookError, "非锚点含真实内容.*A2"):
            reader.read_workbook(path)

    def test_unknown_nonempty_column_requires_structure_review(self):
        def configure(sheet):
            sheet.append(["项目名称", "投标单位名称"])
            sheet.append(["项目甲", "甲公司"])
            sheet["M1"] = "陌生字段"
            sheet["M2"] = False
            sheet["XFD1"].fill = FILL
        book = reader.read_workbook(self.save(configure))
        spec = reader.inspect_workbooks([book])["suggested_plan"]["sources"][0]["sheets"][0]
        self.assertEqual(spec["action"], "needs_mapping")
        self.assertEqual(spec["tables"][0]["column_dispositions"][0]["column"], "M")
        self.assertEqual(spec["tables"][0]["column_dispositions"][0]["disposition"], "unrecognized")

    def test_hidden_sheet_rows_and_column_ranges_keep_content(self):
        def configure(sheet):
            sheet.parent.create_sheet("可见")
            sheet.sheet_state = "hidden"
            sheet.row_dimensions[2].hidden = True
            sheet.column_dimensions.group("J", "L", hidden=True)
            sheet["J2"], sheet["K2"], sheet["L2"] = 0, False, '=""'
            sheet["XFD1"].fill = FILL
        sheet = reader.read_workbook(self.save(configure)).sheets[0]
        self.assertTrue(sheet.hidden)
        self.assertEqual(sheet.hidden_rows, [2])
        self.assertEqual(sheet.hidden_columns, [10, 11, 12])
        self.assertEqual(sheet.raw(2, 10).value, 0)
        self.assertIs(sheet.raw(2, 11).value, False)
        self.assertTrue(sheet.raw(2, 12).formula)

    def test_shared_formulas_dates_and_number_formats_survive(self):
        def configure(sheet):
            sheet.parent.epoch = CALENDAR_MAC_1904
            sheet["A1"], sheet["A2"] = 3, 4
            sheet["B1"], sheet["B2"] = "=A1", "=A2"
            sheet["C1"] = datetime(2025, 7, 18)
            sheet["D1"] = 123
            sheet["D1"].number_format = "00000"
            sheet["E1"] = "#DIV/0!"
            sheet["XFD1"].fill = FILL
        path = self.save(configure)
        def change(root):
            for cell in root.findall(".//" + Q + "c"):
                if cell.get("r") in {"B1", "B2"}:
                    formula = cell.find(Q + "f")
                    formula.attrib.update(t="shared", si="0")
                    if cell.get("r") == "B1":
                        formula.set("ref", "B1:B2")
                    else:
                        formula.text = None
        self.rewrite(path, "xl/worksheets/sheet1.xml", change)
        sheet = reader.read_workbook(path).sheets[0]
        self.assertEqual(sheet.raw(2, 2).value, "=A2")
        self.assertTrue(sheet.raw(2, 2).formula)
        self.assertEqual(sheet.raw(1, 3).value, datetime(2025, 7, 18))
        self.assertEqual(sheet.raw(1, 4).number_format, "00000")
        self.assertTrue(sheet.raw(1, 5).error)

    def test_declared_dimension_cannot_hide_content(self):
        def configure(sheet):
            sheet["A1"] = "首"
            sheet["M5"] = "末"
        path = self.save(configure)
        self.rewrite(path, "xl/worksheets/sheet1.xml", lambda root: root.find(Q + "dimension").set("ref", "A1"))
        self.assertEqual(reader.read_workbook(path).sheets[0].raw(5, 13).value, "末")

    def test_shared_string_zero_index_is_content_but_empty_entry_is_not(self):
        def configure(sheet):
            sheet["A1"] = "占位"
            sheet["XFD1"] = "占位"
        path = self.save(configure)
        def change(root):
            for cell in root.findall(".//" + Q + "c"):
                address = cell.attrib["r"]
                cell.clear()
                cell.attrib.update(r=address, t="s")
                ET.SubElement(cell, Q + "v").text = "0" if address == "A1" else "1"
        self.rewrite(path, "xl/worksheets/sheet1.xml", change)
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        strings = ET.Element(Q + "sst")
        for value in ("真实内容", ""):
            ET.SubElement(ET.SubElement(strings, Q + "si"), Q + "t").text = value
        entries["xl/sharedStrings.xml"] = ET.tostring(strings)
        rels = ET.fromstring(entries["xl/_rels/workbook.xml.rels"])
        ET.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {
            "Id": "sharedStrings", "Target": "/xl/sharedStrings.xml",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"})
        entries["xl/_rels/workbook.xml.rels"] = ET.tostring(rels)
        types = ET.fromstring(entries["[Content_Types].xml"])
        ET.SubElement(types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
            "PartName": "/xl/sharedStrings.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"})
        entries["[Content_Types].xml"] = ET.tostring(types)
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
        sheet = reader.read_workbook(path).sheets[0]
        self.assertEqual(set(sheet.cells), {(1, 1)})
        self.assertEqual(sheet.raw(1, 1).value, "真实内容")
        self.rewrite(path, "xl/worksheets/sheet1.xml", lambda root:
                     setattr(root.find(f".//{Q}c[@r='XFD1']/{Q}v"), "text", "0"))
        with self.assertRaisesRegex(RecoverableWorkbookError, "XFD1"):
            reader.read_workbook(path)

    def test_formula_range_with_remote_results_cannot_be_ignored(self):
        path = self.save(lambda s: s.__setitem__("A1", "=1"))
        self.rewrite(path, "xl/worksheets/sheet1.xml", lambda root:
                     root.find(".//" + Q + "f").attrib.update(t="array", ref="A1:XFD1"))
        with self.assertRaisesRegex(RecoverableWorkbookError, "公式范围=A1:XFD1"):
            reader.read_workbook(path)

    def test_implicit_row_and_cell_coordinates_are_preserved(self):
        def configure(sheet):
            sheet.append(["首", 0])
            sheet.append([False, "末"])
        path = self.save(configure)
        def change(root):
            for element in root.iter():
                if element.tag in {Q + "c", Q + "row"}:
                    element.attrib.pop("r")
        self.rewrite(path, "xl/worksheets/sheet1.xml", change)
        sheet = reader.read_workbook(path).sheets[0]
        self.assertEqual([sheet.raw(r, c).value for r in (1, 2) for c in (1, 2)], ["首", 0, False, "末"])

    def test_many_style_cells_never_expand_read_rectangle(self):
        def configure(sheet):
            sheet["A1"] = "首"
            sheet["M120"] = "末"
            for r in range(1, 121):
                for c in range(16300, 16385):
                    sheet.cell(r, c).fill = FILL
        path = self.save(configure)
        original = ReadOnlyWorksheet.iter_rows
        widths = []
        def bounded(sheet, **kwargs):
            self.assertEqual(kwargs, {"min_row": 1, "max_row": 120, "min_col": 1, "max_col": 13})
            for row in original(sheet, **kwargs):
                widths.append(len(row))
                yield row
        with patch.object(Worksheet, "iter_rows", side_effect=AssertionError("不得使用普通模式矩形扫描")), \
                patch.object(ReadOnlyWorksheet, "iter_rows", bounded):
            sheet = reader.read_workbook(path).sheets[0]
        self.assertEqual(sum(widths), 120 * 13)
        self.assertEqual(len(sheet.cells), 2)

    def test_remote_comments_links_and_table_definitions_are_not_styles(self):
        for kind in ("comment", "hyperlink", "table"):
            with self.subTest(kind=kind):
                def configure(sheet):
                    sheet.append(["第一列", "第二列"])
                    sheet.append(["数据", "数据"])
                    if kind == "comment":
                        sheet["XFD1"].comment = Comment("需要审阅的内容", "test")
                    elif kind == "hyperlink":
                        sheet["XFD1"].hyperlink = "https://example.com/evidence"
                        sheet["XFD1"].value = None
                        sheet["XFD1"].fill = FILL
                    else:
                        sheet.add_table(Table(displayName="Data", ref="A1:B2"))
                path = self.save(configure)
                if kind == "table":
                    self.rewrite(path, "xl/tables/table1.xml", lambda root: root.set("ref", "A1:XFD2"))
                with self.assertRaises(RecoverableWorkbookError):
                    reader.read_workbook(path)

    def test_empty_hyperlink_preserves_normal_reader_value(self):
        def configure(sheet):
            sheet["A1"] = "标题"
            sheet["B2"].hyperlink = "https://example.com/evidence"
            sheet["B2"].value = None
            sheet["B2"].fill = FILL
        sheet = reader.read_workbook(self.save(configure)).sheets[0]
        self.assertEqual(sheet.raw(2, 2).value, "https://example.com/evidence")

    def test_filter_range_does_not_define_content_extent(self):
        def configure(sheet):
            sheet["A1"] = "投标单位名称"
            sheet["A2"] = "甲公司"
            sheet.auto_filter.ref = "A1:XFD2"
            sheet["XFD1"].fill = FILL
        sheet = reader.read_workbook(self.save(configure)).sheets[0]
        self.assertEqual((sheet.max_row, sheet.max_col), (2, 1))

    def test_resource_limits_remain_batch_fatal(self):
        def configure(sheet):
            sheet.append([1, 2])
            sheet.append([3, 4])
            other = sheet.parent.create_sheet("另一表")
            other.append([1, 2])
            other.append([3, 4])
        path = self.save(configure)
        with patch.object(reader, "MAX_CELLS", 6):
            with self.assertRaises(LedgerError) as error:
                reader.read_workbook(path)
            self.assertNotIsInstance(error.exception, RecoverableWorkbookError)
        with patch.object(reader, "MAX_BYTES", 10):
            with self.assertRaises(LedgerError) as error:
                reader.read_workbook(path)
            self.assertNotIsInstance(error.exception, RecoverableWorkbookError)

    def test_sheet_count_is_checked_before_loading_cells(self):
        def configure(sheet):
            for index in range(64):
                sheet.parent.create_sheet(str(index))
        path = self.save(configure)
        with patch.object(openpyxl, "load_workbook") as load:
            with self.assertRaisesRegex(LedgerError, "Sheet 数超限"):
                reader.read_workbook(path)
            load.assert_not_called()

    def test_xml_entities_are_rejected_without_loading_workbook(self):
        path = self.save(lambda s: s.__setitem__("A1", "正常数据"))
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries["xl/worksheets/sheet1.xml"] = (
            b'<!DOCTYPE worksheet [<!ENTITY injected "unexpected">]>' + entries["xl/worksheets/sheet1.xml"])
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
        with patch.object(openpyxl, "load_workbook") as load:
            with self.assertRaisesRegex(RecoverableWorkbookError, "DTDForbidden"):
                reader.read_workbook(path)
            load.assert_not_called()

    def test_zip_entry_and_decompressed_size_limits_remain_fatal(self):
        for kind in ("entries", "size"):
            with self.subTest(kind=kind):
                path = self.save(lambda s: s.__setitem__("A1", "正常数据"))
                with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as archive:
                    if kind == "entries":
                        for index in range(4096):
                            archive.writestr(f"padding/{index}", b"")
                    else:
                        with archive.open("padding", "w") as stream:
                            block = b"\0" * (1024 * 1024)
                            for _ in range(129):
                                stream.write(block)
                with self.assertRaisesRegex(LedgerError, "解压规模超限") as error:
                    reader.read_workbook(path)
                self.assertNotIsInstance(error.exception, RecoverableWorkbookError)

    def test_loader_permission_failure_is_not_a_recoverable_file_error(self):
        path = self.save(lambda s: s.__setitem__("A1", "正常数据"))
        with patch.object(openpyxl, "load_workbook", side_effect=PermissionError("denied")):
            with self.assertRaises(PermissionError):
                reader.read_workbook(path)

    def test_multifile_range_failures_are_audited_in_both_orders(self):
        def normal(sheet):
            sheet.append(["投标单位名称"])
            sheet.append(["甲公司"])
        good = self.save(normal, "A.xlsx")
        for kind in ("value", "formula", "merge"):
            def configure(sheet):
                if kind == "merge":
                    sheet["A1"] = "内容"
                    sheet.merge_cells("A1:KN1")
                else:
                    sheet.cell(1, 300, "内容" if kind == "value" else '=""')
            bad = self.save(configure, "B.xlsx")
            for order in ((good, bad), (bad, good)):
                with self.subTest(kind=kind, first=order[0].name):
                    output = self.directory / f"out-{kind}-{order[0].stem}"
                    first, payload = self.cli("run", *order, "--output", output)
                    self.assertEqual(first.returncode, 3, first.stderr)
                    self.assertFalse(output.exists())
                    inspection = payload["inspection"]
                    failure = inspection["source_failures"][0]
                    self.assertEqual(failure["file_name"], "B.xlsx")
                    self.assertIn("KN1", failure["error"])
                    self.assertEqual(failure["sha256"], hashlib.sha256(bad.read_bytes()).hexdigest())
                    plan = self.directory / "plan.json"
                    plan.write_text(json.dumps(inspection["suggested_plan"]), encoding="utf-8")
                    second, _ = self.cli("run", *order, "--plan", plan, "--output", output)
                    self.assertEqual(second.returncode, 0, second.stderr)
                    ledger = json.loads((output / "ledger.json").read_text())
                    self.assertEqual(ledger["summary"]["record_count"], 1)
                    self.assertEqual(ledger["unique_companies"], ["甲公司"])
                    self.assertEqual([i["code"] for i in ledger["issues"]], ["SOURCE_UNREADABLE"])
                    with (output / "review_queue.csv").open(encoding="utf-8-sig", newline="") as stream:
                        review = list(csv.DictReader(stream))
                    self.assertEqual(len(review), 1)
                    self.assertEqual(review[0]["公司名称"], "")

    def test_resolve_keeps_failed_file_isolated(self):
        def configure(sheet):
            sheet.append(["项目名称", "投标单位名称", "中标单位"])
            sheet.append(["项目甲", "甲建设有限公司", "甲建设工程有限公司"])
        good = self.save(configure, "A.xlsx")
        bad = self.save(lambda s: s.__setitem__("XFD1", "真实内容"), "B.xlsx")
        output = self.directory / "resolved"
        first, payload = self.cli("run", good, bad, "--output", output)
        self.assertEqual(first.returncode, 3, first.stderr)
        plan = self.directory / "plan.json"
        plan.write_text(json.dumps(payload["inspection"]["suggested_plan"]), encoding="utf-8")
        second, payload = self.cli("run", good, bad, "--plan", plan, "--output", output)
        self.assertEqual(second.returncode, 4, second.stderr)
        answers = self.directory / "answers.json"
        answers.write_text(json.dumps({q["question_id"]: q["options"][0]["value"] for q in payload["questions"]}),
                           encoding="utf-8")
        result, _ = self.cli("resolve", "--state", payload["state"], "--answers", answers)
        self.assertEqual(result.returncode, 0, result.stderr)
        ledger = json.loads((output / "ledger.json").read_text())
        self.assertEqual(ledger["summary"]["record_count"], 1)
        self.assertEqual(sum(i["code"] == "SOURCE_UNREADABLE" for i in ledger["issues"]), 1)

    def test_only_failed_sources_publish_audit_without_business_records(self):
        bad = self.save(lambda s: s.__setitem__("XFD1", "真实内容"))
        other = self.save(lambda s: s.__setitem__("XFD1", "另一份内容"), "other.xlsx")
        for inputs in ((bad,), (bad, other)):
            with self.subTest(count=len(inputs)):
                result, inspection = self.cli("inspect", *inputs)
                self.assertEqual(result.returncode, 0, result.stderr)
                self.assertEqual(inspection["sources"], [])
                self.assertEqual(len(inspection["source_failures"]), len(inputs))
                output = self.directory / f"audit-only-{len(inputs)}"
                result, payload = self.cli("run", *inputs, "--output", output)
                self.assertEqual(result.returncode, 0, result.stderr)
                self.assertEqual(payload["kind"], "result")
                ledger = json.loads((output / "ledger.json").read_text())
                self.assertEqual(ledger["summary"]["record_count"], 0)
                self.assertEqual(ledger["unique_companies"], [])
                self.assertTrue(all(s["status"] == "unreadable" for s in ledger["sources"]))
                self.assertEqual([i["code"] for i in ledger["issues"]], ["SOURCE_UNREADABLE"] * len(inputs))


if __name__ == "__main__":
    unittest.main()
