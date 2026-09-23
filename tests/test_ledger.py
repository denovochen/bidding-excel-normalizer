"""验证表格变体与产物业务不变量。 @author denovochen"""
from __future__ import annotations

import copy
import csv
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from ledger_core.artifacts import publish, validate_outputs
from ledger_core.contract import FIELDS, LedgerError
from ledger_core.normalize import aliases_from_json, build_outputs, normalize_name, split_names
from ledger_core.workbook import inspect_workbooks, load_json, read_workbook, validate_plan

TIME = "2026-09-20T12:00:00+08:00"


class LedgerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.aliases = {"schema_version": 1, "aliases": []}

    def tearDown(self):
        self.temp.cleanup()

    def book(self, rows, merges=(), extra_sheet=None):
        w = openpyxl.Workbook()
        s = w.active
        s.title = "台账"
        for row in rows:
            s.append(row)
        for merge in merges:
            s.merge_cells(merge)
        if extra_sheet:
            other = w.create_sheet("补充")
            for row in extra_sheet:
                other.append(row)
        path = self.directory / "input.xlsx"
        w.save(path)
        w.close()
        b = read_workbook(path)
        plan = inspect_workbooks([b])["suggested_plan"]
        for source in plan["sources"]:
            for sheet in source["sheets"]:
                for table in sheet.get("tables", []):
                    table["award_completeness"] = {"status": "complete", "basis_type": "structural",
                                                   "basis": "测试夹具声明为完整最终结果区域"}
                    # 既有名称匹配测试显式使用该模式；逐行模式另用 inspect 的实际建议验证。
                    table["award_mode"] = "name_match"
        return b, plan

    def build(self, book, plan):
        return build_outputs([book], plan, self.aliases, TIME)

    def build_auto(self, book):
        plan = inspect_workbooks([book])["suggested_plan"]
        for source in plan["sources"]:
            for sheet in source["sheets"]:
                for table in sheet.get("tables", []):
                    table["award_completeness"] = {"status": "complete", "basis_type": "structural",
                                                   "basis": "测试夹具声明为完整最终结果区域"}
        return self.build(book, plan)

    def test_group_match_rejects_global_alias_configuration(self):
        b, p = self.book([["项目名称", "投标单位名称", "中标单位"],
                          ["项目甲", "甲建设有限公司", "甲建设工程有限公司"]])
        self.aliases["aliases"] = [{"from": "甲建设有限公司", "to": "甲建设工程有限公司", "basis": "测试全局映射"}]
        with self.assertRaisesRegex(LedgerError, "不接受全局"):
            self.build(b, p)

    def test_row_presence_does_not_assume_first_row_or_only_one_winner(self):
        b, _ = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位"],
                          ["项目甲", "甲公司", 3, None], [None, "乙公司", None, "乙公司"],
                          [None, "丙公司", None, "丙公司"]], ("A2:A4",))
        f, r, ledger = self.build_auto(b)
        self.assertEqual([x["中标与否"] for x in f], ["否", "是", "是"])
        self.assertFalse(r)

    def test_summary_is_not_an_extra_record_and_non_tender_keeps_stated_winner(self):
        b, _ = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位", "备注"],
                          ["项目甲", "项目汇总", 5, "甲公司、乙公司、丙公司"],
                          [None, "甲公司", 2, "甲公司"], [None, "丙公司", None, None],
                          [None, "乙公司", 2, "乙公司"], [None, "丁公司", None, None],
                          [None, "丙公司", 1, "丙公司", "未招投标"]], ("A2:A7",))
        f, r, ledger = self.build_auto(b)
        self.assertEqual([x["中标与否"] for x in f], ["是", "否", "是", "否", "是"])
        self.assertEqual(len(f), 5)
        self.assertFalse(r)
        self.assertEqual(ledger["records"][-1]["participation_type"], "non_tender")
        self.assertEqual(ledger["summary"]["group_count"], 3)
        publish(self.directory / "out", f, r, ledger)

    def test_unmatched_singleton_does_not_win_just_because_award_cell_is_nonempty(self):
        b, _ = self.book([["投标单位名称", "中标单位"], ["甲建设有限公司", "乙建设有限公司"]])
        f, r, ledger = self.build_auto(b)
        self.assertEqual(f[0]["公司名称"], "甲建设有限公司")
        self.assertEqual(f[0]["中标与否"], "")
        self.assertTrue(r)
        self.assertEqual(ledger["groups"][0]["award_matches"][0]["status"], "unresolved")

    def test_row_presence_formula_is_unknown_not_a_win(self):
        b, _ = self.book([["投标单位名称", "中标单位"], ["甲公司", "=A2"], ["乙公司", None]])
        f, r, ledger = self.build_auto(b)
        self.assertEqual([x["中标与否"] for x in f], ["", ""])
        self.assertEqual(len(r), 1)
        self.assertEqual(ledger["issues"][0]["code"], "FIELD_UNAVAILABLE")

    def test_merged_award_uses_group_matching_not_top_row_presence(self):
        b, _ = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位"],
                          ["项目甲", "甲公司", 3, "丙公司"], [None, "乙公司"], [None, "丙公司"]], ("A2:A4", "D2:D4"))
        p = inspect_workbooks([b])["suggested_plan"]
        self.assertEqual(p["sources"][0]["sheets"][0]["tables"][0]["award_mode"], "auto")
        p["sources"][0]["sheets"][0]["tables"][0]["award_completeness"] = {
            "status": "complete", "basis_type": "structural", "basis": "测试夹具声明为完整最终结果区域"}
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["中标与否"] for x in f], ["否", "否", "是"])
        self.assertFalse(r)
        p["sources"][0]["sheets"][0]["tables"][0]["award_mode"] = "row_presence"
        with self.assertRaisesRegex(LedgerError, "已停用"):
            self.build(b, p)

    def test_mixed_structure_beyond_first_twelve_rows_requests_mapping(self):
        rows = [["项目名称", "投标单位名称", "中标单位"]]
        rows.extend([["项目甲", "甲公司", "甲公司"] for _ in range(13)])
        rows.append(["项目甲", "甲公司、乙公司", "甲公司"])
        b, _ = self.book(rows)
        p = inspect_workbooks([b])["suggested_plan"]
        self.assertEqual(p["sources"][0]["sheets"][0]["action"], "needs_mapping")

    def test_mixed_sections_can_reuse_header_and_preserve_project_identity(self):
        b, p = self.book([["项目名称", "标段", "投标单位名称", "中标单位", "投标单位数量"],
                          ["项目甲", "一标", "甲公司", "甲公司", 1],
                          [None, "二标", "乙公司", "丙公司", 2], [None, "二标", "丙公司"],
                          [None, "三标", "丁公司、戊公司", "丁公司", 2],
                          [None, "四标", "己公司", None, 1]], ("A2:A6", "D3:D4"))
        spec = p["sources"][0]["sheets"][0]
        self.assertEqual(spec["action"], "needs_mapping")
        base = {"header_rows": [1], "columns": {"project_name": "A", "lot_name": "B", "bidder_name": "C", "award_name": "D", "bidder_count": "E"},
                "project_mode": "merged", "group_mode": "lot", "award_list_complete": True,
                "summary_markers": ["项目汇总"], "non_tender_markers": []}
        spec.update({"action": "parse", "tables": [
            {**base, "data_start_row": 2, "data_end_row": 2, "award_mode": "name_match", "bidder_separator": "single"},
            {**base, "data_start_row": 3, "data_end_row": 4, "award_mode": "name_match", "bidder_separator": "single"},
            {**base, "data_start_row": 5, "data_end_row": 5, "award_mode": "name_match", "bidder_separator": "delimited"},
            {**base, "data_start_row": 6, "data_end_row": 6, "award_mode": "name_match", "bidder_separator": "single"},
        ]})
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["中标与否"] for x in f], ["是", "否", "是", "是", "否", ""])
        self.assertEqual(ledger["summary"]["project_count"], 1)
        self.assertEqual(ledger["summary"]["group_count"], 4)
        self.assertFalse(r)
        publish(self.directory / "out", f, r, ledger)

    def test_shared_header_does_not_allow_overlapping_data(self):
        b, p = self.book([["项目名称", "投标单位名称", "中标单位"], ["项目甲", "甲公司", "甲公司"], ["项目乙", "乙公司", "乙公司"]])
        spec = p["sources"][0]["sheets"][0]
        spec["tables"].append(copy.deepcopy(spec["tables"][0]))
        with self.assertRaisesRegex(LedgerError, "重叠"):
            self.build(b, p)

    def test_blank_duplicate_row_does_not_conflict_with_shared_group_award(self):
        b, _ = self.book([["项目名称", "投标单位名称", "中标单位"],
                          ["项目甲", "甲公司", "甲公司"], ["项目甲", "甲公司", None]])
        f, r, ledger = self.build_auto(b)
        self.assertEqual(len(f), 1)
        self.assertEqual(f[0]["中标与否"], "是")
        self.assertFalse(r)
        self.assertEqual(len(ledger["records"][0]["occurrences"]), 2)
        publish(self.directory / "out", f, r, ledger)

    def test_row_presence_with_direct_status_conflict_is_not_silently_overwritten(self):
        b, _ = self.book([["投标单位名称", "中标单位", "是否中标"], ["甲公司", "甲公司", "否"]])
        f, r, ledger = self.build_auto(b)
        self.assertEqual(f[0]["中标与否"], "")
        self.assertEqual(len(r), 1)
        self.assertIn("AWARD_STATUS_CONFLICT", [i["code"] for i in ledger["issues"]])

    def test_exact_winner_and_three_artifact_contract(self):
        b, p = self.book([["项目名称", "标段", "投标企业名单", "中标单位"], ["项目甲", "一标", "甲公司、乙公司", "甲公司"]])
        before = hashlib.sha256(b.path.read_bytes()).hexdigest()
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["中标与否"] for x in f], ["是", "否"])
        self.assertEqual(r, [])
        out = self.directory / "output"
        publish(out, f, r, ledger)
        self.assertEqual({x.name for x in out.iterdir()}, {"final.csv", "review_queue.csv", "ledger.json"})
        for name in ("final.csv", "review_queue.csv"):
            self.assertTrue((out / name).read_bytes().startswith(b"\xef\xbb\xbf"))
            with (out / name).open(encoding="utf-8-sig", newline="") as stream:
                self.assertEqual(next(csv.reader(stream)), FIELDS)
        self.assertTrue(validate_outputs(out)["validated"])
        self.assertEqual(hashlib.sha256(b.path.read_bytes()).hexdigest(), before)

    def test_name_matching_preserves_bidder_with_audit(self):
        b, p = self.book([["项目名称", "投标企业名单", "中标单位"],
                          ["项目甲", "甲建设有限公司、乙公司", "甲建设工程有限公司"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual(f[0]["公司名称"], "甲建设有限公司")
        self.assertEqual([x["中标与否"] for x in f], ["", ""])
        self.assertEqual(len(r), 2)
        self.assertEqual(ledger["records"][0]["occurrences"][0]["raw_company"], "甲建设有限公司")
        self.assertEqual(ledger["records"][0]["corrections"], [])
        self.assertEqual(ledger["groups"][0]["award_matches"][0]["status"], "unresolved")
        publish(self.directory / "out", f, r, ledger)

    def test_matching_keeps_bidder_spelling_and_original_award(self):
        b, p = self.book([["项目名称", "投标企业名单", "中标单位"],
                          ["项目甲", "海岳建筑工程有限公司、乙公司、丙公司", "海岳建筑公司有限公司"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["中标与否"] for x in f], ["", "", ""])
        self.assertEqual(len(r), 3)
        self.assertEqual(ledger["groups"][0]["awards"][0]["raw"], "海岳建筑公司有限公司")
        self.assertEqual(ledger["groups"][0]["awards"][0]["name"], "海岳建筑公司有限公司")
        self.assertEqual(ledger["summary"]["record_count"], 3)
        publish(self.directory / "out", f, r, ledger)

    def test_ambiguous_fuzzy_candidates_do_not_become_winners_or_rewrite_names(self):
        b, p = self.book([["项目名称", "投标企业名单", "中标单位"], ["甲项目", "甲建设有限公司、甲建筑有限公司", "甲建设工程有限公司"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["中标与否"] for x in f], ["", ""])
        self.assertEqual(f[0]["公司名称"], "甲建设有限公司")
        self.assertEqual(len(r), 2)
        self.assertEqual(len(ledger["issues"]), 1)
        self.assertGreater(ledger["issues"][0]["candidates"][0]["name_similarity"], .7)
        publish(self.directory / "out", f, r, ledger)

    def test_repeated_prefix_can_match_without_correcting_bidder(self):
        b, p = self.book([["项目名称", "投标企业名单", "中标单位"],
                          ["项目甲", "东州东州市海岳塑料制品有限公司、乙公司", "东州市海岳塑料制品有限公司"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual(f[0]["公司名称"], "东州东州市海岳塑料制品有限公司")
        self.assertEqual(f[0]["中标与否"], "")
        self.assertEqual(len(r), 2)
        self.assertEqual(ledger["records"][0]["corrections"], [])

    def test_repeated_prefix_collision_does_not_merge_two_bidders(self):
        b, p = self.book([["项目名称", "投标企业名单", "中标单位"],
                          ["甲项目", "石家石家庄甲公司、石家庄甲公司", "石家庄甲公司"]])
        f, _, _ = self.build(b, p)
        self.assertEqual(len(f), 2)
        self.assertEqual(f[0]["公司名称"], "石家石家庄甲公司")
        self.assertEqual(f[0]["中标与否"], "否")
        self.assertEqual(f[0]["复核状态"], "通过")

    def test_trailing_annotation_cleanup(self):
        b, p = self.book([["项目名称", "投标企业名单", "中标单位"], ["甲项目", "甲公司投标、乙公司", "甲公司"]])
        f, r, _ = self.build(b, p)
        self.assertEqual(f[0]["公司名称"], "甲公司")
        self.assertFalse(r)

    def test_missing_award_and_rank_one_do_not_infer_award(self):
        b, p = self.book([["项目名称", "投标单位名称", "排名"], ["甲项目", "甲公司", 1]])
        f, r, _ = self.build(b, p)
        self.assertEqual(f[0]["投标排名"], "1")
        self.assertEqual(f[0]["中标与否"], "")
        self.assertEqual(len(r), 0)

    def test_incomplete_award_list_never_marks_losers(self):
        b, p = self.book([["项目名称", "投标企业名单", "中标单位"], ["甲项目", "甲公司、乙公司", "甲公司"]])
        p["sources"][0]["sheets"][0]["tables"][0]["award_completeness"] = {
            "status": "partial", "basis_type": "explicit", "basis": "原表明确说明仅记录部分中标结果"}
        f, _, _ = self.build(b, p)
        self.assertEqual([r["中标与否"] for r in f], ["是", ""])

    def test_non_tender_note_does_not_clear_stated_award(self):
        b, p = self.book([["项目名称", "投标单位名称", "中标单位", "备注"], ["甲项目", "甲公司", "甲公司", "该项目未招投标"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual(f[0]["中标与否"], "是")
        self.assertEqual(len(r), 0)
        self.assertEqual(ledger["records"][0]["participation_type"], "non_tender")

    def test_negated_non_tender_note_does_not_override_explicit_award(self):
        b, p = self.book([["项目名称", "投标企业名单", "中标单位", "备注"],
                          ["甲项目", "甲公司、乙公司", "甲公司", "不存在未招投标情形，已完成公开招标"],
                          ["乙项目", "丙公司", "丙公司", "该项目未招投标"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["中标与否"] for x in f], ["是", "否", "是"])
        self.assertEqual(len(r), 0)
        self.assertEqual(ledger["summary"]["non_tender_group_count"], 1)

    def test_uncertain_procurement_is_not_claimed_as_non_tender(self):
        b, p = self.book([["项目名称", "投标企业名单", "中标单位", "备注"],
                          ["甲项目", "甲公司、乙公司", "甲公司", "是否未招投标，尚待核实"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["中标与否"] for x in f], ["是", "否"])
        self.assertEqual(len(r), 0)
        self.assertFalse(ledger["groups"][0]["non_tender"])
        self.assertEqual(ledger["summary"]["uncertain_group_count"], 1)
        self.assertEqual(ledger["summary"]["bidding_group_count"], 0)
        publish(self.directory / "out", f, r, ledger)

    def test_dedup_only_within_group_preserves_every_occurrence(self):
        b, p = self.book([["项目名称", "标段", "投标企业名单", "中标单位"],
                          ["甲项目", "一标", "甲公司、甲公司、乙公司", "甲公司"],
                          ["甲项目", "二标", "甲公司、乙公司", "乙公司"]])
        f, _, ledger = self.build(b, p)
        self.assertEqual(len(f), 4)
        self.assertEqual(ledger["summary"]["duplicate_mentions_removed"], 1)
        self.assertEqual(len(ledger["records"][0]["occurrences"]), 2)

    def test_duplicate_conflicting_prices_are_preserved_in_review(self):
        b, p = self.book([["项目名称", "投标单位名称", "中标单位", "投标报价"],
                          ["甲项目", "甲公司", "甲公司", 1], ["甲项目", "甲公司", "甲公司", 2]])
        f, r, ledger = self.build(b, p)
        self.assertEqual(len(f), 1)
        self.assertEqual(len(r), 1)
        self.assertIn("DUPLICATE_CONFLICT", [x["code"] for x in ledger["issues"]])
        self.assertEqual(len(ledger["records"][0]["occurrences"]), 2)

    def test_merged_same_name_projects_are_not_collapsed(self):
        b, p = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位"],
                          ["同名项目", "甲公司", 2, "甲公司"], [None, "乙公司"],
                          ["同名项目", "甲公司", 2, "甲公司"], [None, "丙公司"]], ("A2:A3", "A4:A5"))
        f, _, ledger = self.build(b, p)
        self.assertEqual(ledger["summary"]["project_count"], 2)
        self.assertEqual(len(f), 4)
        self.assertNotEqual(ledger["records"][0]["project_id"], ledger["records"][2]["project_id"])

    def test_blocks_inherit_project_but_not_company(self):
        b, p = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位"],
                          ["甲项目", "甲公司", 2, "甲公司"], [None, "乙公司"], [None, None]])
        p["sources"][0]["sheets"][0]["tables"][0]["project_mode"] = "blocks"
        f, _, ledger = self.build(b, p)
        self.assertEqual([r["项目名称"] for r in f], ["甲项目", "甲项目"])
        self.assertEqual(ledger["records"][1]["occurrences"][0]["cells"]["project_name"], "A2")

    def test_summary_and_agent_are_not_enterprises(self):
        b, p = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位", "代理公司"],
                          ["甲项目", "项目汇总", 2, "甲公司", "代理甲公司"],
                          [None, "甲公司", 2, "甲公司", "代理乙公司"], [None, "乙公司"]], ("A2:A4",))
        f, _, ledger = self.build(b, p)
        self.assertEqual(len(f), 2)
        self.assertEqual(ledger["unique_companies"], ["甲公司", "乙公司"])
        self.assertEqual(ledger["summary"]["group_count"], 1)

    def test_count_mismatch_is_reviewed(self):
        b, p = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位"], ["甲项目", "甲公司", 3, "甲公司"]])
        _, r, ledger = self.build(b, p)
        self.assertEqual(len(r), 1)
        self.assertIn("BIDDER_COUNT_MISMATCH", [x["code"] for x in ledger["issues"]])

    def test_award_only_at_group_first_row_applies_to_whole_group(self):
        b, p = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位"],
                          ["同一项目", "甲公司", 2, "甲公司"], [None, "乙公司"],
                          [None, "丙公司", 2, "丙公司"], [None, "丁公司"]], ("A2:A5",))
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["中标与否"] for x in f], ["是", "否", "是", "否"])
        self.assertFalse(r)
        self.assertEqual(ledger["summary"]["project_count"], 1)
        self.assertEqual(ledger["summary"]["group_count"], 2)
        self.assertIn("原中标字段=D2:甲公司", f[1]["证据文本"])
        self.assertIn("原中标字段=D4:丙公司", f[3]["证据文本"])

    def test_missing_award_in_next_group_does_not_inherit_previous_winner(self):
        b, p = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位"],
                          ["同一项目", "甲公司", 2, "甲公司"], [None, "乙公司"],
                          [None, "甲公司", 2, None], [None, "丙公司"]], ("A2:A5",))
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["中标与否"] for x in f], ["是", "否", "", ""])
        self.assertFalse(r)
        self.assertEqual(ledger["groups"][1]["awards"], [])

    def test_first_row_name_typo_not_blank_continuations_causes_review(self):
        b, p = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位"],
                          ["项目甲", "甲建设工程有限公司", 3, "甲建设公司有限公司"],
                          [None, "乙公司"], [None, "丙公司"]], ("A2:A4",))
        f, r, ledger = self.build(b, p)
        self.assertEqual(len(r), 3)
        self.assertEqual([i["code"] for i in ledger["issues"]], ["AWARD_NAME_MISMATCH"])
        self.assertTrue(all("原中标字段=D2:" in x["证据文本"] for x in f))

    def test_reordered_columns_shifted_headers_and_two_sheets(self):
        rows = [["标题"], ["中标企业", "工程名称", "投标人名称"], ["甲公司", "甲项目", "甲公司"]]
        b, p = self.book(rows, extra_sheet=[["项目名称", "投标单位名称", "中标单位"], ["乙项目", "乙公司", "乙公司"]])
        f, _, ledger = self.build(b, p)
        self.assertEqual(len(f), 2)
        self.assertEqual([r["项目名称"] for r in f], ["甲项目", "乙项目"])
        self.assertEqual(ledger["summary"]["project_count"], 2)

    def test_multilevel_header_and_parenthesized_delimiter(self):
        b, p = self.book([["项目基本情况", None, "招投标情况", None],
                          ["项目名称", "标段", "投标企业名单", "中标单位"],
                          ["甲项目", "一标", "甲(上海,中国)公司、乙公司", "甲(上海,中国)公司"]], ("A1:B1", "C1:D1"))
        p["sources"][0]["sheets"][0]["tables"][0]["header_rows"] = [1, 2]
        p["sources"][0]["sheets"][0]["ignored_rows"] = []
        f, _, _ = self.build(b, p)
        self.assertEqual(len(f), 2)
        self.assertEqual(f[0]["公司名称"], "甲(上海,中国)公司")

    def test_line_wrapping_not_a_fake_enterprise(self):
        tokens = split_names("河北甲建设\n工程有限公司", "delimited")
        self.assertEqual(tokens, ["河北甲建设\n工程有限公司"])
        name, actions = normalize_name(tokens[0], {})
        self.assertEqual(name, "河北甲建设工程有限公司")
        self.assertEqual(actions[0]["before"], tokens[0])
        self.assertEqual(split_names("甲公司\n乙公司", "delimited"), ["甲公司", "乙公司"])

    def test_raw_cell_whitespace_and_split_normalization_are_audited(self):
        original = " 甲公司、 乙公司\n"
        b, p = self.book([["项目名称", "投标企业名单", "中标单位"], ["甲项目", original, "甲公司"]])
        f, _, ledger = self.build(b, p)
        self.assertEqual([r["公司名称"] for r in f], ["甲公司", "乙公司"])
        for record in ledger["records"]:
            self.assertEqual(record["occurrences"][0]["raw_values"]["bidder_name"], original)
            self.assertTrue(record["corrections"])
        self.assertEqual(ledger["records"][1]["occurrences"][0]["raw_company"], " 乙公司\n")

    def test_wrong_group_mode_cannot_silently_merge_distinct_lots(self):
        b, p = self.book([["项目名称", "标段", "投标单位名称", "中标单位"],
                          ["甲项目", "一标", "甲公司", "甲公司"], ["甲项目", "二标", "乙公司", "乙公司"]])
        p["sources"][0]["sheets"][0]["tables"][0]["group_mode"] = "project"
        with self.assertRaisesRegex(LedgerError, "标段信息不一致"):
            self.build(b, p)

    def test_consortium_creates_standalone_review_not_fake_bidder(self):
        b, p = self.book([["项目名称", "标段", "投标企业名单", "中标单位"],
                          ["甲项目", "一标", "甲公司、乙公司", "甲公司"],
                          ["乙项目", "二标", "甲公司与乙公司联合体", "甲公司与乙公司联合体"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual(len(f), 2)
        self.assertEqual(len(r), 1)
        self.assertEqual(r[0]["公司名称"], "")
        publish(self.directory / "out", f, r, ledger)

    def test_formula_field_is_blank_and_reviewed_without_blocking(self):
        b, p = self.book([["项目名称", "投标单位名称", "中标单位"], ["甲项目", "=A1", "甲公司"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual(f[0]["公司名称"], "")
        self.assertEqual(f[0]["项目名称"], "甲项目")
        self.assertEqual(len(r), 1)
        self.assertEqual(ledger["unique_companies"], [])
        self.assertIn("FIELD_UNAVAILABLE", [i["code"] for i in ledger["issues"]])
        self.assertEqual(ledger["records"][0]["occurrences"][0]["raw_values"]["bidder_name"], "=A1")
        publish(self.directory / "out", f, r, ledger)

    def test_roster_without_projects_exports_blank_business_fields(self):
        b, p = self.book([["序号", "投标单位名称"], [1, "甲公司"], [2, "甲 公司"], [3, "乙公司"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual(len(f), 2)
        for row in f:
            for key in ("项目名称", "项目编号", "标段名称", "标段编号", "中标与否", "投标排名"):
                self.assertEqual(row[key], "")
        self.assertEqual(r, [])
        self.assertEqual(ledger["projects"], [])
        self.assertEqual(ledger["summary"]["project_count"], 0)
        self.assertEqual(ledger["summary"]["bidding_group_count"], 0)
        self.assertEqual(ledger["summary"]["roster_group_count"], 1)
        self.assertTrue(all(rec["project_id"] is None for rec in ledger["records"]))
        self.assertEqual(ledger["unique_companies"], ["甲公司", "乙公司"])
        publish(self.directory / "out", f, r, ledger)

    def test_only_award_names_are_kept_as_explicit_winners(self):
        b, p = self.book([["中标单位"], ["甲公司"], ["乙公司"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["公司名称"] for x in f], ["甲公司", "乙公司"])
        self.assertEqual([x["中标与否"] for x in f], ["是", "是"])
        self.assertFalse(r)
        self.assertEqual(ledger["projects"], [])
        self.assertTrue(all(x["participation_type"] == "award_company" for x in ledger["records"]))
        publish(self.directory / "out", f, r, ledger)

    def test_project_only_table_keeps_fields_without_fake_company(self):
        b, p = self.book([["项目名称", "项目编号"], ["甲项目", "P001"], ["乙项目", "P002"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["项目名称"] for x in f], ["甲项目", "乙项目"])
        self.assertTrue(all(x["公司名称"] == "" for x in f))
        self.assertFalse(r)
        self.assertEqual(ledger["unique_companies"], [])
        publish(self.directory / "out", f, r, ledger)

    def test_explicit_award_status_is_preserved_without_award_name(self):
        b, p = self.book([["投标单位名称", "是否中标"], ["甲公司", "是"], ["乙公司", "否"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["中标与否"] for x in f], ["是", "否"])
        self.assertFalse(r)
        publish(self.directory / "out", f, r, ledger)

    def test_missing_project_cell_does_not_block_known_company(self):
        b, p = self.book([["项目名称", "投标单位名称"], [None, "甲公司"], ["乙项目", "乙公司"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual([x["项目名称"] for x in f], ["", "乙项目"])
        self.assertEqual([x["公司名称"] for x in f], ["甲公司", "乙公司"])
        self.assertEqual(ledger["summary"]["project_count"], 1)
        self.assertFalse(r)
        publish(self.directory / "out", f, r, ledger)

    def test_project_row_with_empty_company_is_kept_without_fabrication(self):
        b, p = self.book([["项目名称", "投标单位名称", "中标单位", "备注"],
                          ["甲项目", "甲公司", "甲公司", "已完成"], ["乙项目", None, None, "设计状态"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual(len(f), 2)
        self.assertEqual(f[1]["项目名称"], "乙项目")
        self.assertEqual(f[1]["公司名称"], "")
        self.assertEqual(f[1]["中标与否"], "")
        self.assertFalse(r)
        self.assertEqual(ledger["unique_companies"], ["甲公司"])
        self.assertEqual(ledger["summary"]["context_group_count"], 1)
        publish(self.directory / "out", f, r, ledger)

    def test_empty_company_in_existing_group_is_not_marked_a_loser(self):
        b, p = self.book([["项目名称", "投标单位名称", "中标单位"],
                          ["甲项目", "甲公司", "甲公司"], ["甲项目", None, None]])
        f, r, ledger = self.build(b, p)
        self.assertEqual(len(f), 2)
        self.assertEqual(f[1]["中标与否"], "")
        self.assertEqual(ledger["records"][1]["participation_type"], "context")
        publish(self.directory / "out", f, r, ledger)

    def test_conflicting_direct_status_is_not_resolved_by_first_row(self):
        b, p = self.book([["项目名称", "投标单位名称", "是否中标"],
                          ["甲项目", "甲公司", "是"], ["甲项目", "甲公司", "否"]])
        f, r, ledger = self.build(b, p)
        self.assertEqual(len(f), 1)
        self.assertEqual(f[0]["中标与否"], "")
        self.assertEqual(len(r), 1)
        publish(self.directory / "out", f, r, ledger)

    def test_multisource_collection_list_dedups_without_deleting_participation(self):
        b, p = self.book([["项目名称", "投标单位名称"], ["项目甲", "ＡＣＭＥ公司"], ["项目乙", "乙公司"]])
        second = self.directory / "roster.xlsx"
        w = openpyxl.Workbook()
        w.active.append(["投标单位名称"])
        w.active.append(["acme公司"])
        w.active.append(["丙公司"])
        w.save(second); w.close()
        books = [b, read_workbook(second)]
        plan = inspect_workbooks(books)["suggested_plan"]
        f, r, ledger = build_outputs(books, plan, self.aliases, TIME)
        self.assertEqual(len(f), 4)
        self.assertEqual(ledger["unique_companies"], ["ACME公司", "乙公司", "丙公司"])
        self.assertEqual(f[2]["项目名称"], "")
        self.assertFalse(r)
        publish(self.directory / "out", f, r, ledger)

    def test_run_cli_reviews_roster_structure_then_completes_without_question(self):
        b, _ = self.book([["序号", "投标单位名称"], [1, "甲公司"], [2, "甲公司"], [3, "乙公司"]])
        output = self.directory / "direct-output"
        first = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "excel_ledger.py"), "run",
                                str(b.path), "--output", str(output)], capture_output=True, text=True)
        self.assertEqual(first.returncode, 3)
        plan_path = self.directory / "roster-plan.json"
        initial = json.loads(first.stdout)
        stored = json.loads(Path(initial["inspection"]["inspection_path"]).read_text(encoding="utf-8"))
        plan_path.write_text(json.dumps(stored["suggested_plan"], ensure_ascii=False),
                             encoding="utf-8")
        result = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "excel_ledger.py"), "run",
                                 str(b.path), "--plan", str(plan_path), "--output", str(output)],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout)["kind"], "result")
        names = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "excel_ledger.py"), "companies", str(output)],
                                capture_output=True, text=True)
        self.assertEqual(names.returncode, 0, names.stderr)
        self.assertEqual(json.loads(names.stdout)["companies"], ["甲公司", "乙公司"])
        self.assertEqual(len(list(output.iterdir())), 3)

    def test_progress_reports_only_completed_script_stages_before_result(self):
        b, _ = self.book([["投标单位名称"], ["甲公司"]])
        out = self.directory / "progress-output"
        plan_path = self.directory / "progress-plan.json"
        plan_path.write_text(json.dumps(inspect_workbooks([b])["suggested_plan"], ensure_ascii=False), encoding="utf-8")
        result = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "excel_ledger.py"), "run",
                                 str(b.path), "--plan", str(plan_path), "--output", str(out), "--progress"],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        messages = [json.loads(line) for line in result.stdout.splitlines()]
        events = messages[:-1]
        self.assertEqual([(e["step"], e["status"]) for e in events],
                         [(1, "in_progress"), (1, "completed"), (2, "in_progress"), (2, "completed"),
                          (3, "in_progress"), (3, "completed")])
        self.assertTrue(all(e["kind"] == "progress" for e in events))
        self.assertEqual(messages[-1]["kind"], "result")
        self.assertTrue((out / "ledger.json").is_file())
        self.assertFalse(any(e.get("step") == 4 for e in messages))

    def test_progress_failure_does_not_claim_success_or_delivery(self):
        result = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "excel_ledger.py"), "run",
                                 str(self.directory / "missing.xlsx"), "--progress"], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        events = [json.loads(line) for line in result.stdout.splitlines()]
        self.assertEqual([(e["step"], e["status"]) for e in events], [(1, "in_progress"), (1, "failed")])
        self.assertEqual(json.loads(result.stderr)["kind"], "error")

    def test_mapping_required_keeps_cleaning_open_and_no_artifacts(self):
        b, _ = self.book([["自定义组织字段", "说明"], ["甲公司", "示例"]])
        out = self.directory / "unknown-output"
        result = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "excel_ledger.py"), "run",
                                 str(b.path), "--output", str(out), "--progress"], capture_output=True, text=True)
        self.assertEqual(result.returncode, 3)
        messages = [json.loads(line) for line in result.stdout.splitlines()]
        self.assertEqual(messages[-1]["kind"], "mapping_required")
        self.assertEqual([(e["step"], e["status"]) for e in messages[:-1]],
                         [(1, "in_progress"), (1, "completed"), (2, "in_progress")])
        self.assertFalse(out.exists())

    def test_official_code_zero_format_is_preserved(self):
        b, p = self.book([["项目名称", "项目编号", "投标单位名称", "中标单位"], ["甲项目", 42, "甲公司", "甲公司"]])
        w = openpyxl.load_workbook(b.path)
        w.active["B2"].number_format = "00000"
        w.save(b.path); w.close()
        b = read_workbook(b.path)
        p["sources"][0]["sha256"] = b.sha256
        f, _, _ = self.build(b, p)
        self.assertEqual(f[0]["项目编号"], "00042")

    def test_source_hash_drift_rejected(self):
        b, p = self.book([["项目名称", "投标单位名称"], ["甲项目", "甲公司"]])
        p["sources"][0]["sha256"] = "0" * 64
        with self.assertRaisesRegex(LedgerError, "哈希"):
            self.build(b, p)

    def test_unmapped_sheet_and_uncovered_rows_rejected(self):
        b, p = self.book([["项目名称", "投标单位名称"], ["甲项目", "甲公司"], ["乙项目", "乙公司"]])
        bad = copy.deepcopy(p); bad["sources"][0]["sheets"] = []
        with self.assertRaises(LedgerError):
            self.build(b, bad)
        p["sources"][0]["sheets"][0]["tables"][0]["data_end_row"] = 2
        with self.assertRaisesRegex(LedgerError, "未解释"):
            self.build(b, p)

    def test_unknown_mapping_code_is_rejected(self):
        b, p = self.book([["项目名称", "投标单位名称"], ["甲项目", "甲公司"]])
        p["sources"][0]["sheets"][0]["tables"][0]["python"] = "print(123)"
        with self.assertRaises(LedgerError):
            self.build(b, p)

    def test_multiple_tables_and_repeated_header_require_explicit_scope(self):
        b, p = self.book([["项目名称", "投标单位名称", "中标单位"], ["甲项目", "甲公司", "甲公司"],
                          ["项目名称", "投标单位名称", "中标单位"], ["乙项目", "乙公司", "乙公司"]])
        spec = p["sources"][0]["sheets"][0]
        self.assertEqual(spec["action"], "needs_mapping")
        spec["action"] = "parse"
        template = {"columns": {"project_name": "A", "bidder_name": "B", "award_name": "C"},
                    "project_mode": "repeated", "group_mode": "project", "bidder_separator": "single",
                    "award_list_complete": True, "summary_markers": ["合计"], "non_tender_markers": ["未招投标"]}
        first = {**template, "header_rows": [1], "data_start_row": 2, "data_end_row": 4}
        spec["tables"] = [first]; spec["ignored_rows"] = []
        with self.assertRaisesRegex(LedgerError, "表头"):
            self.build(b, p)
        first["data_end_row"] = 2
        second = {**template, "header_rows": [3], "data_start_row": 4, "data_end_row": 4}
        spec["tables"] = [first, second]
        f, _, _ = self.build(b, p)
        self.assertEqual(len(f), 2)

    def test_ids_and_rows_repeat_deterministically(self):
        b, p = self.book([["项目名称", "投标企业名单", "中标单位"], ["甲项目", "甲公司、乙公司", "甲公司"]])
        a = self.build(b, p)
        other = build_outputs([b], p, self.aliases, "2026-09-21T12:00:00+08:00")
        self.assertEqual([r["id"] for r in a[2]["records"]], [r["id"] for r in other[2]["records"]])
        self.assertEqual(a[2]["groups"], other[2]["groups"])

    def test_tampering_and_existing_output_are_rejected(self):
        b, p = self.book([["项目名称", "投标单位名称", "中标单位"], ["甲项目", "甲公司", "甲公司"]])
        f, r, ledger = self.build(b, p)
        out = self.directory / "out"
        publish(out, f, r, ledger)
        with self.assertRaisesRegex(LedgerError, "覆盖"):
            publish(out, f, r, ledger)
        with (out / "final.csv").open("ab") as stream:
            stream.write(b"tampered")
        with self.assertRaisesRegex(LedgerError, "哈希"):
            validate_outputs(out)

    def test_csv_dangerous_literal_is_escaped_and_audited(self):
        b, p = self.book([["项目名称", "投标单位名称", "中标单位"], ["@项目", "+甲公司", "+甲公司"]])
        f, r, ledger = self.build(b, p)
        out = self.directory / "out"
        publish(out, f, r, ledger)
        self.assertIn("'+甲公司", (out / "final.csv").read_text(encoding="utf-8-sig"))
        self.assertTrue(validate_outputs(out)["validated"])

    def test_malformed_and_fake_extension_are_rejected(self):
        path = self.directory / "fake.xls"
        path.write_bytes(b"not an excel workbook")
        with self.assertRaises(LedgerError):
            read_workbook(path)

    def test_alias_cycle_and_duplicate_rule_are_rejected(self):
        payload = {"schema_version": 1, "aliases": [{"from": "甲公司", "to": "乙公司", "basis": "用户确认"},
                                                    {"from": "乙公司", "to": "甲公司", "basis": "用户确认"}]}
        with self.assertRaises(LedgerError):
            aliases_from_json(payload)

    def test_cli_error_has_nonzero_exit_and_no_success_files(self):
        result = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "excel_ledger.py"), "export",
                                 str(self.directory / "missing.xlsx"), "--plan", str(self.directory / "missing.json"),
                                 "--output", str(self.directory / "out")], capture_output=True, text=True)
        self.assertEqual(result.returncode, 2)
        self.assertEqual(json.loads(result.stderr)["kind"], "error")
        self.assertFalse((self.directory / "out").exists())

    def test_unknown_bidder_column_requires_review_and_is_not_silently_dropped(self):
        b, plan = self.book([["项目名称", "参标单位", "中标单位"], ["项目甲", "甲公司", "甲公司"]])
        output = self.directory / "unknown-bidder-output"
        entry = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "excel_ledger.py"), "run",
                                str(b.path), "--output", str(output)], capture_output=True, text=True)
        self.assertEqual(entry.returncode, 3)
        self.assertFalse(output.exists())
        spec = plan["sources"][0]["sheets"][0]
        self.assertEqual(spec["action"], "needs_mapping")
        table = spec["tables"][0]
        self.assertNotIn("bidder_name", table["columns"])
        self.assertEqual(next(item for item in table["column_dispositions"] if item["column"] == "B")["disposition"],
                         "unrecognized")
        with self.assertRaisesRegex(LedgerError, "结构审阅"):
            self.build(b, plan)
        table["columns"]["bidder_name"] = "B"
        table["column_dispositions"] = [item for item in table["column_dispositions"] if item["column"] != "B"]
        table["structure_warnings"] = []
        spec["action"] = "parse"
        final, review, _ = self.build(b, plan)
        self.assertEqual([(row["公司名称"], row["中标与否"]) for row in final], [("甲公司", "是")])
        self.assertFalse(review)

    def test_unmerged_project_value_at_block_start_is_inherited(self):
        b, plan = self.book([["项目名称", "投标单位名称", "中标单位"],
                             ["项目甲", "甲公司", "甲公司"], [None, "乙公司", None]])
        table = plan["sources"][0]["sheets"][0]["tables"][0]
        self.assertEqual(table["project_mode"], "blocks")
        final, _, ledger = self.build(b, plan)
        self.assertEqual([row["项目名称"] for row in final], ["项目甲", "项目甲"])
        self.assertEqual(ledger["summary"]["project_count"], 1)

    def test_repeated_bidder_count_is_not_used_as_group_anchor(self):
        b, plan = self.book([["项目名称", "投标单位名称", "投标单位数量", "中标单位"],
                             ["项目甲", "甲公司", 3, "甲公司"], ["项目甲", "乙公司", 3, None],
                             ["项目甲", "丙公司", 3, None]])
        table = plan["sources"][0]["sheets"][0]["tables"][0]
        self.assertEqual(table["group_mode"], "project")
        final, review, ledger = self.build(b, plan)
        self.assertEqual(ledger["summary"]["group_count"], 1)
        self.assertEqual([row["中标与否"] for row in final], ["是", "否", "否"])
        self.assertFalse(review)

    def test_group_context_separates_same_project_lot_across_batches(self):
        b, plan = self.book([["项目名称", "批次", "标段", "投标单位名称", "中标单位"],
                             ["项目甲", "第一批", "一标", "甲公司", "甲公司"],
                             ["项目甲", "第二批", "一标", "甲公司", "甲公司"]])
        spec = plan["sources"][0]["sheets"][0]
        self.assertEqual(spec["action"], "needs_mapping")
        table = spec["tables"][0]
        item = next(item for item in table["column_dispositions"] if item["column"] == "B")
        item.update({"disposition": "group_context", "mode": "repeated", "reason": "批次构成独立招标组边界"})
        table["structure_warnings"] = []
        spec["action"] = "parse"
        final, review, ledger = self.build(b, plan)
        self.assertEqual(len(final), 2)
        self.assertEqual(ledger["summary"]["group_count"], 2)
        self.assertEqual([row["中标与否"] for row in final], ["是", "是"])
        self.assertFalse(review)

    def test_sparse_complete_award_column_can_mark_non_winners(self):
        b, plan = self.book([["项目名称", "投标单位名称", "中标单位"],
                             ["项目甲", "甲公司", "甲公司"], ["项目甲", "乙公司", None]])
        final, review, ledger = self.build(b, plan)
        self.assertEqual([row["中标与否"] for row in final], ["是", "否"])
        self.assertTrue(ledger["groups"][0]["award_completeness"]["verified"])
        self.assertFalse(review)

    def test_late_header_is_found_before_execution(self):
        rows = [["项目名称", "投标单位名称", "中标单位"]]
        rows.extend([["项目甲", f"企业{index}公司", None] for index in range(12)])
        rows.extend([["项目名称", "投标单位名称", "中标单位"], ["项目乙", "乙公司", "乙公司"]])
        b, plan = self.book(rows)
        spec = plan["sources"][0]["sheets"][0]
        self.assertEqual(spec["action"], "needs_mapping")
        warnings = spec["tables"][0]["structure_warnings"]
        self.assertIn(14, [warning.get("row") for warning in warnings if warning["code"] == "NEW_HEADER"])

    def test_merged_award_status_does_not_spread_to_multiple_bidders(self):
        b, plan = self.book([["项目名称", "投标单位名称", "是否中标"],
                             ["项目甲", "甲公司", "是"], ["项目甲", "乙公司", None]], ("C2:C3",))
        final, review, ledger = self.build(b, plan)
        self.assertEqual([row["中标与否"] for row in final], ["", ""])
        self.assertEqual([issue["code"] for issue in ledger["issues"]], ["AWARD_STATUS_SCOPE_AMBIGUOUS"])
        self.assertEqual(len(review), 2)

    def test_duplicate_nonempty_fields_merge_independently_of_row_order(self):
        results = []
        for detail_rows in (
            [["项目甲", "甲公司", None, None], ["项目甲", "甲公司", "是", 1]],
            [["项目甲", "甲公司", "是", 1], ["项目甲", "甲公司", None, None]],
        ):
            b, plan = self.book([["项目名称", "投标单位名称", "是否中标", "排名"], *detail_rows])
            final, review, ledger = self.build(b, plan)
            results.append((final[0]["中标与否"], final[0]["投标排名"], len(review), ledger["summary"]["issue_count"]))
        self.assertEqual(results, [("是", "1", 0, 0), ("是", "1", 0, 0)])

    def test_multiple_unavailable_fields_on_one_row_keep_distinct_issues(self):
        b, plan = self.book([["项目名称", "投标单位名称", "中标单位"], ["项目甲", "=A2", "=A2"]])
        _, _, ledger = self.build(b, plan)
        field_issues = [issue for issue in ledger["issues"] if issue["code"] == "FIELD_UNAVAILABLE"]
        self.assertEqual({issue["field"] for issue in field_issues}, {"bidder_name", "award_name"})
        self.assertEqual(len({issue["id"] for issue in field_issues}), 2)

    def test_review_region_delivers_other_rows_without_fake_record(self):
        b, plan = self.book([["项目名称", "投标单位名称"], ["项目甲", "甲公司"], ["项目乙", "乙公司"]])
        spec = plan["sources"][0]["sheets"][0]
        spec["tables"][0]["data_end_row"] = 2
        spec["review_regions"] = [{"start": 3, "end": 3, "reason": "局部结构仍无法确认"}]
        final, review, ledger = self.build(b, plan)
        self.assertEqual([row["公司名称"] for row in final], ["甲公司"])
        self.assertEqual(len(review), 1)
        self.assertEqual(review[0]["公司名称"], "")
        self.assertEqual(ledger["summary"]["record_count"], 1)
        self.assertEqual(ledger["summary"]["unique_company_count"], 1)

    def test_multifile_unreadable_source_keeps_valid_delivery(self):
        b, _ = self.book([["项目名称", "投标单位名称"], ["项目甲", "甲公司"]])
        bad = self.directory / "broken.xlsx"
        bad.write_bytes(b"not-ooxml")
        output = self.directory / "mixed-output"
        first = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "excel_ledger.py"), "run",
                                str(b.path), str(bad), "--output", str(output)], capture_output=True, text=True)
        self.assertEqual(first.returncode, 3, first.stderr)
        summary = json.loads(first.stdout)["inspection"]
        inspection = json.loads(Path(summary["inspection_path"]).read_text(encoding="utf-8"))
        self.assertEqual(len(inspection["source_failures"]), 1)
        plan_path = self.directory / "mixed-plan.json"
        plan_path.write_text(json.dumps(inspection["suggested_plan"], ensure_ascii=False), encoding="utf-8")
        second = subprocess.run([sys.executable, "-B", str(ROOT / "scripts" / "excel_ledger.py"), "run",
                                 str(b.path), str(bad), "--plan", str(plan_path), "--output", str(output)],
                                capture_output=True, text=True)
        self.assertEqual(second.returncode, 0, second.stderr)
        ledger = json.loads((output / "ledger.json").read_text(encoding="utf-8"))
        self.assertEqual(ledger["summary"]["record_count"], 1)
        self.assertEqual(ledger["summary"]["unique_company_count"], 1)
        self.assertIn("SOURCE_UNREADABLE", [issue["code"] for issue in ledger["issues"]])
        with (output / "review_queue.csv").open(encoding="utf-8-sig", newline="") as stream:
            review = list(csv.DictReader(stream))
        self.assertEqual(len(review), 1)
        self.assertEqual(review[0]["公司名称"], "")

    def test_confidence_is_rule_based_and_audited(self):
        b, plan = self.book([["项目名称", "投标单位名称"], ["项目甲", "甲公司"]])
        final, _, ledger = self.build(b, plan)
        self.assertNotEqual(final[0]["置信度"], "1.0")
        self.assertEqual(float(final[0]["置信度"]), ledger["records"][0]["confidence"]["score"])
        self.assertEqual(ledger["records"][0]["confidence"]["meaning"], "rule_reliability_not_probability")


if __name__ == "__main__":
    unittest.main()
