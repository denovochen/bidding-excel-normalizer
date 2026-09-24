"""平台短响应、批次隔离、幂等恢复和有界证据的回归。 @author denovochen"""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import excel_agent as agent
from ledger_core.contract import LedgerError
from ledger_core.workbook import load_json


class AgentProtocolTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def book(self, wide=False):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["项目名称", "投标企业名称"] + ([f"附加信息{i}" for i in range(45)] if wide else []))
        for index in range(25):
            sheet.append(["测试项目", f"企业{index}有限公司"] + (["来源说明" * 15] * 45 if wide else []))
        path = self.root / "含 空格输入.xlsx"
        workbook.save(path)
        workbook.close()
        return path

    def run_start(self, wide=False, budget=6000):
        return agent.start([self.book(wide)], self.root / "result", budget)

    def answers(self, reply):
        path = Path(reply["answer_file"])
        payload = load_json(path)
        for answer in payload["answers"].values():
            answer["basis"] = "表为逐企业投标记录；附加信息为辅助来源"
            answer["columns"] = {key: "evidence" if value == "unresolved" else value
                                 for key, value in answer["columns"].items()}
        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        return path

    def test_cli_normal_handoff_is_success_and_stdout_is_one_json(self):
        process = subprocess.run([sys.executable, str(ROOT / "scripts/excel_agent.py"), "run",
                                  str(self.book()), "--output", str(self.root / "out")],
                                 capture_output=True, text=True, cwd=self.root, timeout=30)
        self.assertEqual(process.returncode, 0, process.stderr)
        reply = json.loads(process.stdout)
        self.assertEqual(reply["kind"], "mapping_required")
        self.assertLessEqual(len(process.stdout), 6000)
        self.assertEqual(process.stderr, "")
        self.assertTrue(Path(reply["answer_file"]).exists())
        self.assertIn("answer_roles", reply)

    def test_invalid_cli_arguments_also_return_json_without_traceback(self):
        process = subprocess.run([sys.executable, str(ROOT / "scripts/excel_agent.py"), "inspect",
                                  "--state", "missing", "--region", "region", "--start", "not-a-row"],
                                 capture_output=True, text=True, timeout=30)
        self.assertEqual(process.returncode, 2)
        self.assertEqual(json.loads(process.stdout)["kind"], "error")
        self.assertEqual(process.stderr, "")

    def test_submit_replay_returns_same_result_and_conflicting_replay_fails(self):
        reply = self.run_start()
        state = Path(reply["state"])
        answers = self.answers(reply)
        result = agent.submit(state, answers)
        self.assertEqual(result["kind"], "result")
        self.assertEqual(result["summary"]["record_count"], 25)
        self.assertEqual(agent.submit(state, answers), result)
        payload = load_json(answers)
        answers.write_text(json.dumps(payload, sort_keys=True), encoding="utf-8")
        self.assertEqual(agent.submit(state, answers), result)
        self.assertEqual({p.name for p in (self.root / "result").iterdir()}, {"final.csv", "review_queue.csv", "ledger.json"})
        payload = load_json(answers)
        payload["answers"][next(iter(payload["answers"]))]["basis"] = "不同决定"
        answers.write_text(json.dumps(payload), encoding="utf-8")
        with self.assertRaisesRegex(LedgerError, "不同答案"):
            agent.submit(state, answers)

    def test_stale_or_foreign_questions_fail_without_mutating_state(self):
        reply = self.run_start()
        state = Path(reply["state"])
        answers = self.answers(reply)
        original = load_json(answers)
        before = state.read_bytes()
        for payload in ({**original, "state_version": 0},
                        {**original, "answers": {"previous_question": {}}}):
            answers.write_text(json.dumps(payload), encoding="utf-8")
            with self.assertRaises(LedgerError):
                agent.submit(state, answers)
            self.assertEqual(state.read_bytes(), before)

    def test_wide_question_pages_preserve_all_column_evidence(self):
        reply = self.run_start(wide=True, budget=4000)
        self.assertLessEqual(len(agent.encoded(reply)), 4000)
        question = reply["questions"][0]
        self.assertTrue(question["details_required"])
        offset, columns = 0, []
        while True:
            page = agent.question_page(Path(reply["state"]), question["question_id"], "columns", offset, 8)
            self.assertLessEqual(len(agent.encoded(page)), 4000)
            self.assertTrue(page["items"])
            columns.extend(page["items"])
            offset += len(page["items"])
            if page["next_command"] is None:
                break
        self.assertEqual(len(columns), 47)
        self.assertEqual(len({column["column"] for column in columns}), 47)

    def test_inspect_uses_count_instead_of_inclusive_end_arithmetic(self):
        reply = self.run_start()
        state = Path(reply["state"])
        full = load_json(state)["reply"]["questions"][0]
        result = agent.inspect_page(state, full["regions"][0]["id"], 2, 20, None)
        self.assertEqual(result["returned_row_count"], 20)
        self.assertIn("A21", result["rows"][-1])
        self.assertLessEqual(len(agent.encoded(result)), 6000)

    def test_delivery_counts_come_from_issues_not_invented_optional_fields(self):
        from ledger_core.artifacts import delivery_summary
        ledger = {"summary": {"record_count": 40, "review_record_count": 35, "issue_count": 1},
                  "issues": [{"code": "AWARD_NAME_MISMATCH", "message": "中标名称未确认",
                              "review_sequences": list(range(1, 36))}],
                  "audit_warnings": [], "resolutions": [{"decision": "deferred"}]}
        result = delivery_summary(ledger)
        self.assertEqual(result["audit_warning_count"], 0)
        self.assertEqual(result["review_reasons"][0]["affected_review_records"], 35)
        self.assertEqual(result["review_reasons"][0]["issue_count"], 1)
        self.assertNotIn("法人", result["message"])

    def test_relation_selection_requires_rationale_and_both_source_references(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(["项目名称", "中标单位"])
        workbook.active.append(["清溪项目（财政补助）", "甲公司"])
        sheet = workbook.create_sheet("名册")
        sheet.append(["项目名称", "投标单位名称"])
        sheet.append(["清溪项目", "甲公司"])
        source = self.root / "跨表.xlsx"
        workbook.save(source)
        workbook.close()
        reply = agent.start([source], self.root / "relation-output", 6000)
        while reply["kind"] == "mapping_required":
            path = Path(reply["answer_file"])
            payload = load_json(path)
            for question in load_json(Path(reply["state"]))["reply"]["questions"]:
                key = question["question_id"]
                if key not in payload["answers"]:
                    continue
                if question["task_type"] == "structure":
                    payload["answers"][key]["basis"] = "原表列名明确区分汇总和名册"
                else:
                    payload["answers"][key] = {"bidder_sets": [item["set_id"] for item in question["candidate_bidder_sets"]],
                                                "basis": "汇总与名册同属施工项目"}
            path.write_text(json.dumps(payload), encoding="utf-8")
            reply = agent.submit(Path(reply["state"]), path)
        self.assertEqual(reply["kind"], "relationship_review_required")
        question = reply["questions"][0]
        selected = question["options"][0]["value"]
        answer_path = Path(reply["answer_file"])
        payload = load_json(answer_path)
        state_path = Path(reply["state"])
        before = state_path.read_bytes()
        for answer in (selected, {"value": selected, "basis": "同一项目", "evidence": ["伪造来源"]}):
            payload["answers"] = {question["question_id"]: answer}
            answer_path.write_text(json.dumps(payload), encoding="utf-8")
            with self.assertRaises(LedgerError):
                agent.submit(state_path, answer_path)
            self.assertEqual(before, state_path.read_bytes())
        payload["answers"] = {question["question_id"]: {"value": selected, "basis": "仅缺资金来源后缀，主体名称一致",
                                                        "evidence": question["evidence_refs"][selected]}}
        answer_path.write_text(json.dumps(payload), encoding="utf-8")
        result = agent.submit(state_path, answer_path)
        self.assertEqual(result["kind"], "result")
        ledger = load_json(Path(result["output"]) / "ledger.json")
        resolution = ledger["relationship_resolutions"][0]
        self.assertEqual(resolution["actor"], "model")
        self.assertEqual(resolution["evidence_refs"], question["evidence_refs"][selected])


if __name__ == "__main__":
    unittest.main()
